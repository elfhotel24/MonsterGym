import sqlite3
import os
import sys
from datetime import datetime, timedelta
import customtkinter as ctk
from tkinter import messagebox, Toplevel, filedialog
from impresora import ImpresoraTermica

# =============================================================================
# IMPORTS OPCIONALES (Excel, PDF, Calendario)
# =============================================================================
try:
    from tkcalendar import DateEntry
    TIENE_TKCALENDAR = True
except ImportError:
    TIENE_TKCALENDAR = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    TIENE_OPENPYXL = True
except ImportError:
    TIENE_OPENPYXL = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors as rl_colors
    from reportlab.lib.units import cm
    TIENE_REPORTLAB = True
except ImportError:
    TIENE_REPORTLAB = False


# =============================================================================
# FUNCIÓN PARA RUTAS DE RECURSOS (funciona en .py y en .exe)
# =============================================================================
def resource_path(relative_path):
    """Obtiene la ruta absoluta del recurso, funciona en desarrollo y con PyInstaller."""
    try:
        base_path = sys._MEIPASS  # Carpeta temporal al compilar con --onefile
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)


# =============================================================================
# INICIALIZACIÓN DE LA BASE DE DATOS
# =============================================================================
def init_db():
    if getattr(sys, 'frozen', False):
        base_dir = os.path.dirname(sys.executable)
    else:
        base_dir = os.path.dirname(os.path.abspath(__file__))

    db_path = os.path.join(base_dir, "monster_gym.db")
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL,
        rol TEXT NOT NULL
    )""")

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS clientes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        codigo TEXT UNIQUE,
        nombre TEXT NOT NULL,
        telefono TEXT,
        fecha_inscripcion TEXT NOT NULL,
        fecha_vencimiento TEXT
    )""")

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS productos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        codigo TEXT UNIQUE,
        nombre TEXT NOT NULL,
        categoria TEXT NOT NULL,
        precio_venta REAL NOT NULL,
        costo REAL DEFAULT 0,
        stock INTEGER DEFAULT 0
    )""")

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS compras (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        producto_info TEXT NOT NULL,
        cantidad INTEGER NOT NULL,
        costo_producto REAL NOT NULL,
        fecha_compra TEXT NOT NULL
    )""")

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS cajas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        usuario TEXT NOT NULL,
        monto_apertura REAL NOT NULL,
        monto_cierre REAL,
        fecha_apertura TEXT NOT NULL,
        fecha_cierre TEXT,
        estado TEXT DEFAULT 'Abierta'
    )""")

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS facturas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        codigo_factura TEXT UNIQUE,
        caja_id INTEGER,
        cliente_nombre TEXT NOT NULL,
        total REAL NOT NULL,
        fecha TEXT NOT NULL,
        metodo_pago TEXT DEFAULT 'Efectivo',
        FOREIGN KEY(caja_id) REFERENCES cajas(id)
    )""")

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS factura_detalles (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        factura_id INTEGER,
        nombre_producto TEXT NOT NULL,
        cantidad INTEGER NOT NULL,
        precio_unitario REAL NOT NULL,
        FOREIGN KEY(factura_id) REFERENCES facturas(id)
    )""")

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS historial_inventario (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        producto_id INTEGER,
        nombre_producto TEXT NOT NULL,
        tipo_movimiento TEXT NOT NULL,
        cantidad INTEGER NOT NULL,
        fecha TEXT NOT NULL,
        usuario TEXT,
        motivo TEXT
    )""")

    cursor.execute("SELECT COUNT(*) FROM usuarios")
    if cursor.fetchone()[0] == 0:
        cursor.execute("INSERT INTO usuarios (username, password, rol) VALUES ('admin', '123', 'Admin')")
        cursor.execute("INSERT INTO usuarios (username, password, rol) VALUES ('cajero', '123', 'Caja')")
        cursor.execute("INSERT INTO clientes (codigo, nombre, telefono, fecha_inscripcion) VALUES ('CLI-0000', 'Consumidor Final', '000-000-0000', '2026-01-01')")
        cursor.execute("INSERT INTO productos (codigo, nombre, categoria, precio_venta, costo, stock) VALUES ('PROD-0001', 'Proteina Monster Whey', 'Suplementos', 2500.0, 1800.0, 10)")
        cursor.execute("INSERT INTO productos (codigo, nombre, categoria, precio_venta, costo, stock) VALUES ('PROD-0002', 'Agua Dasani', 'Bebidas', 50.0, 25.0, 30)")

    conn.commit()
    conn.close()
    print(f"✅ Base de datos lista en: {db_path}")


init_db()

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("green")


class MonsterGymApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("MONSTER GYM - Sistema POS")
        self.geometry("1280x800")
        self.minsize(1024, 700)

        try:
            self.state("zoomed")
        except Exception:
            try:
                self.attributes("-zoomed", True)
            except Exception:
                pass

        try:
            icon_path = resource_path("monster.ico")
            if os.path.exists(icon_path):
                self.iconbitmap(icon_path)
                print(f"✅ Ícono cargado: {icon_path}")
            else:
                print(f"⚠️ Ícono no encontrado en: {icon_path}")
        except Exception as e:
            print(f"⚠️ No se pudo cargar el ícono: {e}")

        try:
            self.impresora = ImpresoraTermica(nombre="POS-80-Series")
        except Exception as e:
            messagebox.showwarning("Impresora", f"No se pudo conectar con la impresora:\n{e}\n\nLos tickets se imprimirán en consola.")
            self.impresora = None

        self.after(150, self._maximizar_ventana)

        self.usuario_actual = None
        self.rol_actual = None
        self.id_caja_activa = None
        self.carrito_actual = []
        self.workspace = None

        self.contenedor_principal = ctk.CTkFrame(self, fg_color="#0a0a0f")
        self.contenedor_principal.pack(fill="both", expand=True)

        self.mostrar_pantalla_login()

    def _maximizar_ventana(self):
        try:
            self.state("zoomed")
        except Exception:
            try:
                self.attributes("-zoomed", True)
            except Exception:
                pass

    def conectar_db(self):
        if getattr(sys, 'frozen', False):
            base_dir = os.path.dirname(sys.executable)
        else:
            base_dir = os.path.dirname(os.path.abspath(__file__))
        db_path = os.path.join(base_dir, "monster_gym.db")
        return sqlite3.connect(db_path)

    def limpiar_pantalla(self):
        for widget in self.contenedor_principal.winfo_children():
            widget.destroy()
        self.workspace = None

    def crear_cabecera(self, titulo, subtitulo):
        if self.workspace is None:
            return
        for w in self.workspace.winfo_children():
            w.destroy()
        header_frame = ctk.CTkFrame(self.workspace, fg_color="transparent")
        header_frame.pack(fill="x", pady=(5, 10))
        lbl_t = ctk.CTkLabel(header_frame, text=titulo, font=ctk.CTkFont(size=28, weight="bold"), text_color="white")
        lbl_t.pack(anchor="w")
        lbl_s = ctk.CTkLabel(header_frame, text=subtitulo, font=ctk.CTkFont(size=13), text_color="#a8e600")
        lbl_s.pack(anchor="w")

    # =========================================================================
    # INDICADOR DE ESTADO DE CAJA (sidebar)
    # =========================================================================
    def _actualizar_indicador_caja(self):
        """Actualiza el widget del sidebar que muestra si la caja está abierta o cerrada."""
        try:
            if not hasattr(self, 'lbl_estado_caja_sidebar') or self.lbl_estado_caja_sidebar is None:
                return
            if not self.lbl_estado_caja_sidebar.winfo_exists():
                return
            if self.id_caja_activa:
                self.lbl_estado_caja_sidebar.configure(text="🟢 CAJA ABIERTA", text_color="#a8e600")
                self.frame_estado_caja.configure(fg_color="#0e2e0e", border_color="#a8e600")
            else:
                self.lbl_estado_caja_sidebar.configure(text="🔴 CAJA CERRADA", text_color="#ff595e")
                self.frame_estado_caja.configure(fg_color="#2e0e0e", border_color="#ff595e")
        except Exception:
            pass

    # =========================================================================
    # SELECTOR DE FECHA (calendario si tkcalendar está disponible)
    # =========================================================================
    def crear_selector_fecha(self, parent, width=120):
        """Devuelve un widget con calendario (DateEntry) o un CTkEntry normal si no hay tkcalendar."""
        if TIENE_TKCALENDAR:
            try:
                d = DateEntry(
                    parent,
                    width=12,
                    date_pattern="yyyy-mm-dd",
                    background="#a8e600",
                    foreground="black",
                    borderwidth=1,
                    selectbackground="#a8e600",
                    selectforeground="black",
                    normalbackground="#1a1a26",
                    normalforeground="white",
                    weekendbackground="#1a1a26",
                    weekendforeground="#a8e600",
                )
                d.delete(0, "end")  # Empezar vacío
                return d
            except Exception:
                pass
        return ctk.CTkEntry(parent, width=width, placeholder_text="YYYY-MM-DD", font=ctk.CTkFont(size=13))

    # =========================================================================
    # EXPORTAR A EXCEL
    # =========================================================================
    def exportar_reporte_excel(self, columnas, datos, titulo, nombre_archivo_default):
        if not TIENE_OPENPYXL:
            messagebox.showerror(
                "Librería faltante",
                "Para exportar a Excel necesitas instalar openpyxl.\n\nAbre una terminal y ejecuta:\n\npip install openpyxl"
            )
            return
        if not datos:
            messagebox.showwarning("Sin datos", "No hay datos para exportar.")
            return
        ruta = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=nombre_archivo_default
        )
        if not ruta:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Reporte"

            # Título
            ws.append([titulo])
            ws.append([])
            ws.append(list(columnas))

            for fila in datos:
                ws.append([str(v) if v is not None else "" for v in fila])

            # Estilos
            ws["A1"].font = Font(size=14, bold=True, color="1A73E8")
            for i, _ in enumerate(columnas, 1):
                cell = ws.cell(row=3, column=i)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="222222")

            # Ancho de columnas
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for c in col:
                    if c.value is not None:
                        max_len = max(max_len, len(str(c.value)))
                ws.column_dimensions[col_letter].width = min(max_len + 3, 45)

            wb.save(ruta)
            messagebox.showinfo("Éxito", f"Reporte exportado a:\n{ruta}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo exportar: {e}")

    # =========================================================================
    # EXPORTAR A PDF
    # =========================================================================
    def exportar_reporte_pdf(self, columnas, datos, titulo, nombre_archivo_default):
        if not TIENE_REPORTLAB:
            messagebox.showerror(
                "Librería faltante",
                "Para exportar a PDF necesitas instalar reportlab.\n\nAbre una terminal y ejecuta:\n\npip install reportlab"
            )
            return
        if not datos:
            messagebox.showwarning("Sin datos", "No hay datos para exportar.")
            return
        ruta = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF", "*.pdf")],
            initialfile=nombre_archivo_default
        )
        if not ruta:
            return
        try:
            doc = SimpleDocTemplate(
                ruta, pagesize=A4,
                leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                topMargin=1.5 * cm, bottomMargin=1.5 * cm
            )
            styles = getSampleStyleSheet()
            estilo_titulo = ParagraphStyle(
                "TituloCustom", parent=styles["Title"],
                fontSize=16, textColor=rl_colors.HexColor("#1A73E8")
            )
            estilo_sub = ParagraphStyle(
                "SubCustom", parent=styles["Normal"],
                fontSize=9, textColor=rl_colors.HexColor("#666666")
            )

            elementos = []
            elementos.append(Paragraph("MONSTER GYM", estilo_titulo))
            elementos.append(Paragraph(titulo, ParagraphStyle(
                "Sub", parent=styles["Normal"], fontSize=12,
                textColor=rl_colors.HexColor("#333333"), spaceAfter=6
            )))
            elementos.append(Paragraph(
                f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')} · Usuario: {self.usuario_actual}",
                estilo_sub
            ))
            elementos.append(Spacer(1, 0.5 * cm))

            data_tabla = [list(columnas)]
            for fila in datos:
                data_tabla.append([str(v) if v is not None else "" for v in fila])

            tabla = Table(data_tabla, repeatRows=1)
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#222222")),
                ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 10),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("TOPPADDING", (0, 0), (-1, 0), 8),
                ("BACKGROUND", (0, 1), (-1, -1), rl_colors.HexColor("#f5f5f5")),
                ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.HexColor("#888888")),
                ("FONTSIZE", (0, 1), (-1, -1), 9),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            elementos.append(tabla)
            doc.build(elementos)
            messagebox.showinfo("Éxito", f"Reporte exportado a:\n{ruta}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo exportar: {e}")

    # =========================================================================
    # LOGIN
    # =========================================================================
    def mostrar_pantalla_login(self):
        self.limpiar_pantalla()
        login_frame = ctk.CTkFrame(
            self.contenedor_principal, width=480, height=560,
            fg_color="#1a1a2e", corner_radius=20,
            border_width=2, border_color="#a8e600"
        )
        login_frame.place(relx=0.5, rely=0.5, anchor="center")
        login_frame.pack_propagate(False)

        ctk.CTkLabel(login_frame, text="🏋️ MONSTER GYM", font=ctk.CTkFont(size=38, weight="bold"), text_color="#a8e600").pack(pady=(40, 5))
        ctk.CTkLabel(login_frame, text="SISTEMA DE PUNTO DE VENTA", font=ctk.CTkFont(size=14, weight="bold"), text_color="#888899").pack(pady=(0, 30))

        ctk.CTkLabel(login_frame, text="👤 USUARIO", font=ctk.CTkFont(size=12, weight="bold"), text_color="#ccccdd").pack(anchor="w", padx=50, pady=(0, 2))
        self.txt_usuario = ctk.CTkEntry(login_frame, width=340, height=44, placeholder_text="Ingresa tu usuario", font=ctk.CTkFont(size=14), corner_radius=8, border_color="#3a3a5a", fg_color="#12121e")
        self.txt_usuario.pack(pady=(0, 15))
        self.txt_usuario.focus()

        ctk.CTkLabel(login_frame, text="🔒 CONTRASEÑA", font=ctk.CTkFont(size=12, weight="bold"), text_color="#ccccdd").pack(anchor="w", padx=50, pady=(0, 2))
        self.txt_pass = ctk.CTkEntry(login_frame, width=340, height=44, placeholder_text="••••••••", show="*", font=ctk.CTkFont(size=14), corner_radius=8, border_color="#3a3a5a", fg_color="#12121e")
        self.txt_pass.pack(pady=(0, 25))

        ctk.CTkButton(login_frame, text="▶  INICIAR SESIÓN", width=340, height=50, font=ctk.CTkFont(weight="bold", size=15), fg_color="#a8e600", text_color="#0a0a0f", hover_color="#8fcb00", corner_radius=10, command=self.procesar_login).pack(pady=(0, 20))
        ctk.CTkLabel(login_frame, text="© 2026 Monster Gym · Versión 2.0", font=ctk.CTkFont(size=11), text_color="#555566").pack(pady=(10, 20))

        self.txt_usuario.bind("<Return>", lambda e: self.procesar_login())
        self.txt_pass.bind("<Return>", lambda e: self.procesar_login())

    def procesar_login(self):
        user = self.txt_usuario.get().strip()
        pwd = self.txt_pass.get().strip()
        if not user or not pwd:
            messagebox.showerror("Error", "Por favor, ingresa usuario y contraseña.")
            return

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT rol FROM usuarios WHERE username=? AND password=?", (user, pwd))
        resultado = cursor.fetchone()

        if resultado:
            self.usuario_actual = user
            self.rol_actual = resultado[0]

            cursor.execute("SELECT id FROM cajas WHERE usuario=? AND estado='Abierta'", (user,))
            caja_res = cursor.fetchone()

            if caja_res:
                self.id_caja_activa = caja_res[0]
            else:
                self.id_caja_activa = None

            conn.close()
            self.construir_layout_base()
        else:
            conn.close()
            messagebox.showerror("Error", "Credenciales incorrectas.")

    # =========================================================================
    # LAYOUT PRINCIPAL
    # =========================================================================
    def construir_layout_base(self):
        self.limpiar_pantalla()
        self.contenedor_principal.grid_columnconfigure(0, weight=0)
        self.contenedor_principal.grid_columnconfigure(1, weight=1)
        self.contenedor_principal.grid_rowconfigure(0, weight=1)

        sidebar = ctk.CTkFrame(self.contenedor_principal, width=270, fg_color="#111118", corner_radius=0)
        sidebar.grid(row=0, column=0, sticky="nsew")
        sidebar.grid_rowconfigure(13, weight=1)

        ctk.CTkLabel(sidebar, text="MONSTER GYM", font=ctk.CTkFont(size=24, weight="bold"), text_color="#a8e600").grid(row=0, column=0, padx=20, pady=(20, 2), sticky="w")
        ctk.CTkLabel(sidebar, text=f"Operador: {self.usuario_actual}\nRol: {self.rol_actual}", font=ctk.CTkFont(size=13), text_color="gray", justify="left").grid(row=1, column=0, padx=20, pady=(0, 15), sticky="w")

        # ===== INDICADOR DE ESTADO DE CAJA (arriba del sidebar) =====
        self.frame_estado_caja = ctk.CTkFrame(sidebar, fg_color="#2e0e0e", corner_radius=8, border_width=1, border_color="#ff595e")
        self.frame_estado_caja.grid(row=2, column=0, padx=15, pady=(0, 15), sticky="ew")
        self.lbl_estado_caja_sidebar = ctk.CTkLabel(
            self.frame_estado_caja, text="",
            font=ctk.CTkFont(size=13, weight="bold"), text_color="#ff595e"
        )
        self.lbl_estado_caja_sidebar.pack(pady=8, padx=10)
        self._actualizar_indicador_caja()

        modulos = [
            ("📊 Dashboard", self.mostrar_dashboard, ["Admin", "Caja"]),
            ("📅 Suscripciones", self.mostrar_pantalla_suscripciones, ["Admin", "Caja"]),
            ("🧾 Facturación", self.mostrar_pantalla_facturacion, ["Admin", "Caja"]),
            ("💰 Corte de Caja", self.mostrar_pantalla_corte, ["Admin", "Caja"]),
            ("📁 Control de Cajas", self.mostrar_pantalla_control_caja, ["Admin", "Caja"]),
            ("👤 Clientes", self.mostrar_pantalla_clientes, ["Admin", "Caja"]),
            ("📦 Productos", self.mostrar_pantalla_productos, ["Admin"]),
            ("📥 Inventario", self.mostrar_pantalla_inventario, ["Admin"]),
            ("👥 Usuarios", self.mostrar_pantalla_usuarios, ["Admin"]),
            ("📈 Reportes", self.mostrar_pantalla_reportes, ["Admin"]),
        ]

        fila = 3  # Cambia de 2 a 3 por el indicador de caja
        for nombre, funcion, roles_permitidos in modulos:
            if self.rol_actual in roles_permitidos:
                ctk.CTkButton(sidebar, text=nombre, anchor="w", fg_color="transparent", text_color="white", hover_color="#1f1f2a", height=44, font=ctk.CTkFont(size=14), command=funcion).grid(row=fila, column=0, padx=12, pady=3, sticky="ew")
            fila += 1

        ctk.CTkButton(sidebar, text="🚪 Cerrar Sesión", fg_color="#1a1a26", text_color="white", hover_color="#b32121", height=40, font=ctk.CTkFont(size=14), command=self.mostrar_pantalla_login).grid(row=14, column=0, padx=15, pady=20, sticky="ew")

        self.workspace = ctk.CTkFrame(self.contenedor_principal, fg_color="#0a0a0f")
        self.workspace.grid(row=0, column=1, sticky="nsew", padx=25, pady=25)
        self.mostrar_dashboard()

    # =========================================================================
    # DASHBOARD
    # =========================================================================
    def mostrar_dashboard(self):
        self.crear_cabecera("Dashboard", "Resumen operativo en tiempo real")
        main_frame = ctk.CTkFrame(self.workspace, fg_color="transparent")
        main_frame.pack(fill="both", expand=True)

        cards_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        cards_frame.pack(fill="x", pady=(10, 20))
        cards_frame.grid_columnconfigure((0, 1, 2, 3), weight=1, uniform="cards")

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM clientes WHERE codigo != 'CLI-0000'")
        total_clientes = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM productos")
        total_productos = cursor.fetchone()[0]
        today = datetime.now().strftime("%Y-%m-%d")
        cursor.execute("SELECT COALESCE(SUM(total),0) FROM facturas WHERE fecha LIKE ? AND caja_id=?", (today + '%', self.id_caja_activa if self.id_caja_activa else -1))
        ventas_hoy = cursor.fetchone()[0]
        estado_caja = "Abierta" if self.id_caja_activa else "Cerrada"
        fecha_limite = (datetime.now() + timedelta(days=5)).strftime("%Y-%m-%d")
        cursor.execute("SELECT codigo, nombre, fecha_vencimiento FROM clientes WHERE fecha_vencimiento IS NOT NULL AND fecha_vencimiento <= ? AND codigo != 'CLI-0000' ORDER BY fecha_vencimiento ASC", (fecha_limite,))
        alertas_sub = cursor.fetchall()
        conn.close()

        def crear_tarjeta(parent, titulo, valor, color="#a8e600"):
            card = ctk.CTkFrame(parent, fg_color="#14141e", corner_radius=14, height=110)
            card.pack(side="left", fill="both", expand=True, padx=10, pady=5)
            card.pack_propagate(False)
            ctk.CTkLabel(card, text=titulo, font=ctk.CTkFont(size=13), text_color="gray").pack(pady=(15, 2))
            ctk.CTkLabel(card, text=str(valor), font=ctk.CTkFont(size=32, weight="bold"), text_color=color).pack()

        crear_tarjeta(cards_frame, "Clientes Registrados", total_clientes)
        crear_tarjeta(cards_frame, "Productos en Catálogo", total_productos)
        crear_tarjeta(cards_frame, "Ventas del Día", f"RD${ventas_hoy:,.2f}")
        crear_tarjeta(cards_frame, "Estado de Caja", estado_caja, color="#ff595e" if estado_caja == "Cerrada" else "#a8e600")

        split_lists = ctk.CTkFrame(main_frame, fg_color="transparent")
        split_lists.pack(fill="both", expand=True, pady=5)
        split_lists.grid_columnconfigure(0, weight=1)
        split_lists.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(split_lists, text="⚠️ Suscripciones por Vencer / Vencidas", font=ctk.CTkFont(size=15, weight="bold"), text_color="#ff595e").grid(row=0, column=0, sticky="w", padx=10, pady=(10, 5))
        list_alertas = ctk.CTkScrollableFrame(split_lists, fg_color="#14141e")
        list_alertas.grid(row=1, column=0, sticky="nsew", padx=10)

        if alertas_sub:
            for c, n, f in alertas_sub:
                item = ctk.CTkFrame(list_alertas, fg_color="#1e1e2a")
                item.pack(fill="x", pady=2, padx=5)
                estado = "Vencido" if f < today else "Próximo"
                color = "#ff595e" if estado == "Vencido" else "#ffca3a"
                ctk.CTkLabel(item, text=f"{c} | {n} | Vence: {f}", font=ctk.CTkFont(size=13)).pack(side="left", padx=12)
                ctk.CTkLabel(item, text=estado, font=ctk.CTkFont(size=12, weight="bold"), text_color=color).pack(side="right", padx=12)
        else:
            ctk.CTkLabel(list_alertas, text="No hay alertas pendientes.", text_color="gray").pack(pady=20)

        ctk.CTkLabel(split_lists, text="Últimos Clientes Registrados", font=ctk.CTkFont(size=15, weight="bold"), text_color="gray").grid(row=0, column=1, sticky="w", padx=10, pady=(10, 5))
        list_frame = ctk.CTkScrollableFrame(split_lists, fg_color="#14141e")
        list_frame.grid(row=1, column=1, sticky="nsew", padx=10)

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT codigo, nombre, telefono, fecha_inscripcion FROM clientes WHERE codigo != 'CLI-0000' ORDER BY id DESC LIMIT 5")
        ultimos = cursor.fetchall()
        conn.close()
        if ultimos:
            for c, n, t, f in ultimos:
                item = ctk.CTkFrame(list_frame, fg_color="#1e1e2a")
                item.pack(fill="x", pady=2, padx=5)
                ctk.CTkLabel(item, text=f"{c} | {n} | Tel: {t} | Desde: {f}", font=ctk.CTkFont(size=13)).pack(side="left", padx=12)
        else:
            ctk.CTkLabel(list_frame, text="Aún no hay clientes adicionales.", text_color="gray").pack(pady=20)

    # =========================================================================
    # SUSCRIPCIONES
    # =========================================================================
    def mostrar_pantalla_suscripciones(self):
        self.crear_cabecera("Suscripciones", "Gestión de membresías y renovaciones")
        main_frame = ctk.CTkFrame(self.workspace, fg_color="transparent")
        main_frame.pack(fill="both", expand=True)

        filtro_frame = ctk.CTkFrame(main_frame, fg_color="#1a1a26", corner_radius=12)
        filtro_frame.pack(fill="x", padx=8, pady=5)
        ctk.CTkLabel(filtro_frame, text="Buscar:", font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
        entry_buscar = ctk.CTkEntry(filtro_frame, width=200, placeholder_text="Nombre o código")
        entry_buscar.pack(side="left", padx=5)
        ctk.CTkLabel(filtro_frame, text="Estado:", font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
        cmb_estado = ctk.CTkOptionMenu(filtro_frame, values=["Todos", "Activo", "Vencido", "Próximo a vencer (5 días)"], width=180)
        cmb_estado.set("Todos")
        cmb_estado.pack(side="left", padx=5)

        def aplicar_filtros():
            self.cargar_suscripciones(entry_buscar.get().strip(), cmb_estado.get(), list_frame)

        ctk.CTkButton(filtro_frame, text="Filtrar", fg_color="#a8e600", text_color="black", command=aplicar_filtros).pack(side="left", padx=10)

        list_frame = ctk.CTkScrollableFrame(main_frame, fg_color="#14141e", corner_radius=12)
        list_frame.pack(fill="both", expand=True, padx=8, pady=8)
        self.cargar_suscripciones("", "Todos", list_frame)

    def cargar_suscripciones(self, texto_buscar, estado_filtro, list_frame):
        for widget in list_frame.winfo_children():
            widget.destroy()

        header = ctk.CTkFrame(list_frame, fg_color="#1e1e2a")
        header.pack(fill="x", padx=10, pady=(5, 5))
        ctk.CTkLabel(header, text="CÓDIGO", width=80, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="NOMBRE", width=180, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="INSCRIPCIÓN", width=110, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="VENCIMIENTO", width=110, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="ESTADO", width=150, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="ACCIONES", width=130, font=ctk.CTkFont(weight="bold")).pack(side="right", padx=15)

        conn = self.conectar_db()
        cursor = conn.cursor()
        query = "SELECT id, codigo, nombre, fecha_inscripcion, fecha_vencimiento FROM clientes WHERE codigo != 'CLI-0000'"
        params = []
        if texto_buscar:
            query += " AND (codigo LIKE ? OR nombre LIKE ?)"
            params.extend([f"%{texto_buscar}%", f"%{texto_buscar}%"])
        query += " ORDER BY nombre ASC"
        cursor.execute(query, params)
        clientes_sub = cursor.fetchall()
        conn.close()

        hoy = datetime.now().date()
        for id_cli, cod, nom, f_insc, f_vence in clientes_sub:
            diferencia = None
            estado_texto = "Sin suscripción"
            color_estado = "gray"
            if f_vence:
                try:
                    vence = datetime.strptime(f_vence, "%Y-%m-%d").date()
                    diferencia = (vence - hoy).days
                    if diferencia < 0:
                        estado_texto = f"Vencido: {abs(diferencia)} día(s)"
                        color_estado = "#ff595e"
                    elif diferencia == 0:
                        estado_texto = "Vence HOY"
                        color_estado = "#ffca3a"
                    elif diferencia <= 5:
                        estado_texto = f"Vence en {diferencia} días"
                        color_estado = "#ffca3a"
                    else:
                        estado_texto = "Activo"
                        color_estado = "#a8e600"
                except ValueError:
                    estado_texto = "Fecha inválida"
                    color_estado = "#ff595e"

            if estado_filtro != "Todos":
                if estado_filtro == "Activo" and not (diferencia is not None and diferencia > 5): continue
                if estado_filtro == "Vencido" and not (diferencia is not None and diferencia < 0): continue
                if estado_filtro == "Próximo a vencer (5 días)" and not (diferencia is not None and 0 <= diferencia <= 5): continue

            item = ctk.CTkFrame(list_frame, fg_color="#0d0d14")
            item.pack(fill="x", pady=3, padx=10)
            ctk.CTkLabel(item, text=cod, width=80).pack(side="left", padx=5)
            ctk.CTkLabel(item, text=nom, width=180, anchor="w").pack(side="left", padx=5)
            ctk.CTkLabel(item, text=f_insc or "N/A", width=110).pack(side="left", padx=5)
            ctk.CTkLabel(item, text=f_vence or "N/A", width=110).pack(side="left", padx=5)
            ctk.CTkLabel(item, text=estado_texto, text_color=color_estado, width=150, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
            ctk.CTkButton(item, text="Renovar Plan", width=115, height=30, fg_color="#a8e600", text_color="black",
                          command=lambda i=id_cli, n=nom, f=f_vence: self.abrir_modal_renovacion(i, n, f)).pack(side="right", padx=10)

        if not clientes_sub:
            ctk.CTkLabel(list_frame, text="No se encontraron clientes con los filtros aplicados.", text_color="gray").pack(pady=30)

    def abrir_modal_renovacion(self, cliente_id, nombre, f_vence_actual):
        if not self.id_caja_activa:
            messagebox.showerror("Caja Cerrada", "Debes abrir un turno (caja) para cobrar suscripciones.")
            return

        PRECIO_RENOVACION = 1300.00
        modal = ctk.CTkToplevel(self)
        modal.title(f"Renovar Suscripción: {nombre}")
        modal.geometry("450x400")
        modal.grab_set()

        ctk.CTkLabel(modal, text=f"Atleta: {nombre}", font=ctk.CTkFont(size=18, weight="bold"), text_color="#a8e600").pack(pady=(20, 10))
        ctk.CTkLabel(modal, text=f"Vencimiento actual: {f_vence_actual}" if f_vence_actual else "No tiene suscripción activa", font=ctk.CTkFont(size=13), text_color="gray").pack(pady=(0, 20))
        ctk.CTkLabel(modal, text="Precio de Renovación:", font=ctk.CTkFont(weight="bold")).pack(anchor="w", padx=40)
        ctk.CTkLabel(modal, text=f"RD$ {PRECIO_RENOVACION:,.2f}", font=ctk.CTkFont(size=20, weight="bold"), text_color="#a8e600").pack(pady=5)
        ctk.CTkLabel(modal, text="Método de pago:", font=ctk.CTkFont(weight="bold")).pack(anchor="w", padx=40, pady=(10, 0))
        cmb_pago = ctk.CTkOptionMenu(modal, values=["Efectivo", "Transferencia"], width=370, height=40)
        cmb_pago.pack(pady=10)

        def procesar_renovacion():
            metodo = cmb_pago.get()
            hoy = datetime.now()
            if f_vence_actual:
                f_actual_dt = datetime.strptime(f_vence_actual, "%Y-%m-%d")
                nueva_fecha = f_actual_dt + timedelta(days=30) if f_actual_dt > hoy else hoy + timedelta(days=30)
            else:
                nueva_fecha = hoy + timedelta(days=30)

            nueva_fecha_str = nueva_fecha.strftime("%Y-%m-%d")
            fecha_venta_str = hoy.strftime("%Y-%m-%d %H:%M:%S")
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                conn.execute("BEGIN")
                cursor.execute("UPDATE clientes SET fecha_vencimiento = ? WHERE id = ?", (nueva_fecha_str, cliente_id))
                cursor.execute("INSERT INTO facturas (caja_id, cliente_nombre, total, fecha, metodo_pago) VALUES (?, ?, ?, ?, ?)", (self.id_caja_activa, nombre, PRECIO_RENOVACION, fecha_venta_str, metodo))
                fact_id = cursor.lastrowid
                cod_fact = f"FACT-{fact_id:05d}"
                cursor.execute("UPDATE facturas SET codigo_factura = ? WHERE id = ?", (cod_fact, fact_id))
                cursor.execute("INSERT INTO factura_detalles (factura_id, nombre_producto, cantidad, precio_unitario) VALUES (?, ?, 1, ?)", (fact_id, "Renovación de membresía (30 días)", PRECIO_RENOVACION))
                conn.commit()
                messagebox.showinfo("Éxito", f"Suscripción renovada.\nNueva fecha: {nueva_fecha_str}")
                modal.destroy()
                self.mostrar_pantalla_suscripciones()
            except Exception as e:
                conn.rollback()
                messagebox.showerror("Error", f"Fallo al procesar: {e}")
            finally:
                conn.close()

        ctk.CTkButton(modal, text="Cobrar y Renovar", fg_color="#a8e600", text_color="black", height=45, font=ctk.CTkFont(weight="bold", size=14), command=procesar_renovacion).pack(pady=25)

    # =========================================================================
    # CLIENTES
    # =========================================================================
    def mostrar_pantalla_clientes(self):
        self.crear_cabecera("Clientes", "Gestión de atletas")
        main_split = ctk.CTkFrame(self.workspace, fg_color="transparent")
        main_split.pack(fill="both", expand=True)
        main_split.grid_columnconfigure(0, weight=1)
        main_split.grid_columnconfigure(1, weight=2)
        main_split.grid_rowconfigure(0, weight=1)

        form_frame = ctk.CTkFrame(main_split, fg_color="#14141e", corner_radius=12)
        form_frame.grid(row=0, column=0, padx=8, sticky="nsew")
        ctk.CTkLabel(form_frame, text="Nuevo Atleta", font=ctk.CTkFont(size=16, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(20, 10))
        txt_nombre = ctk.CTkEntry(form_frame, placeholder_text="Nombre Completo", height=40, font=ctk.CTkFont(size=14))
        txt_nombre.pack(fill="x", padx=20, pady=6)
        txt_telefono = ctk.CTkEntry(form_frame, placeholder_text="Teléfono", height=40, font=ctk.CTkFont(size=14))
        txt_telefono.pack(fill="x", padx=20, pady=6)
        ctk.CTkLabel(form_frame, text="Fecha de Inscripción (YYYY-MM-DD):", font=ctk.CTkFont(size=12), text_color="gray").pack(anchor="w", padx=20, pady=(10, 2))
        txt_fecha = ctk.CTkEntry(form_frame, placeholder_text="Ej: 2026-01-15", height=40, font=ctk.CTkFont(size=14))
        hoy = datetime.now().strftime("%Y-%m-%d")
        txt_fecha.insert(0, hoy)
        txt_fecha.pack(fill="x", padx=20, pady=6)
        lbl_codigo_generado = ctk.CTkLabel(form_frame, text="", text_color="#a8e600", font=ctk.CTkFont(size=13))
        lbl_codigo_generado.pack(pady=2)

        def acc_guardar_cliente():
            nom = txt_nombre.get().strip()
            tel = txt_telefono.get().strip()
            fecha_str = txt_fecha.get().strip()
            if not nom:
                messagebox.showerror("Error", "El nombre es obligatorio.")
                return
            if fecha_str:
                try:
                    datetime.strptime(fecha_str, "%Y-%m-%d")
                    fecha_insc = fecha_str
                except ValueError:
                    messagebox.showerror("Error", "Formato de fecha inválido. Use YYYY-MM-DD.")
                    return
            else:
                fecha_insc = datetime.now().strftime("%Y-%m-%d")
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                conn.execute("BEGIN")
                cursor.execute("INSERT INTO clientes (nombre, telefono, fecha_inscripcion, fecha_vencimiento) VALUES (?, ?, ?, NULL)", (nom, tel, fecha_insc))
                last_id = cursor.lastrowid
                codigo_auto = f"CLI-{last_id:04d}"
                cursor.execute("UPDATE clientes SET codigo=? WHERE id=?", (codigo_auto, last_id))
                conn.commit()
                messagebox.showinfo("Éxito", f"Cliente creado.\nCódigo: {codigo_auto}\nFecha: {fecha_insc}")
                txt_nombre.delete(0, "end")
                txt_telefono.delete(0, "end")
                txt_fecha.delete(0, "end")
                txt_fecha.insert(0, hoy)
                lbl_codigo_generado.configure(text=f"Último código: {codigo_auto}")
            except Exception as e:
                conn.rollback()
                messagebox.showerror("Error", f"No se pudo guardar: {e}")
            finally:
                conn.close()
                self.mostrar_pantalla_clientes()

        ctk.CTkButton(form_frame, text="Guardar Atleta", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=14), height=42, command=acc_guardar_cliente).pack(fill="x", padx=20, pady=20)

        list_frame = ctk.CTkFrame(main_split, fg_color="#14141e", corner_radius=12)
        list_frame.grid(row=0, column=1, padx=8, sticky="nsew")

        filtro_clientes = ctk.CTkFrame(list_frame, fg_color="transparent")
        filtro_clientes.pack(fill="x", padx=12, pady=(10, 5))
        ctk.CTkLabel(filtro_clientes, text="Buscar cliente:", font=ctk.CTkFont(size=13)).pack(side="left", padx=(0, 8))
        txt_filtro_clientes = ctk.CTkEntry(filtro_clientes, placeholder_text="Nombre, código o teléfono...", height=36, font=ctk.CTkFont(size=13))
        txt_filtro_clientes.pack(side="left", fill="x", expand=True)

        scroll_clientes = ctk.CTkScrollableFrame(list_frame, fg_color="transparent")
        scroll_clientes.pack(fill="both", expand=True, padx=5, pady=5)

        def cargar_clientes_filtro(texto=""):
            for widget in scroll_clientes.winfo_children():
                widget.destroy()
            texto = texto.strip().lower()

            header = ctk.CTkFrame(scroll_clientes, fg_color="#1e1e2a")
            header.pack(fill="x", padx=10, pady=(5, 2))
            ctk.CTkLabel(header, text="CÓDIGO", width=90, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
            ctk.CTkLabel(header, text="NOMBRE", width=190, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
            ctk.CTkLabel(header, text="TELÉFONO", width=120, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
            ctk.CTkLabel(header, text="INSCRIPCIÓN", width=120, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
            ctk.CTkLabel(header, text="ACCIONES", width=145, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)

            conn = self.conectar_db()
            cursor = conn.cursor()
            cursor.execute("SELECT id, codigo, nombre, telefono, fecha_inscripcion FROM clientes WHERE codigo != 'CLI-0000' ORDER BY id DESC")
            items = cursor.fetchall()
            conn.close()

            encontrados = [item for item in items if not texto or any(texto in str(valor or "").lower() for valor in item[1:5])]

            if encontrados:
                for id_, codigo, nombre, telefono, fecha in encontrados:
                    item = ctk.CTkFrame(scroll_clientes, fg_color="#0d0d14")
                    item.pack(fill="x", pady=2, padx=10)
                    ctk.CTkLabel(item, text=codigo, width=90).pack(side="left", padx=5)
                    ctk.CTkLabel(item, text=nombre, width=190, anchor="w").pack(side="left", padx=5)
                    ctk.CTkLabel(item, text=telefono or "N/A", width=120).pack(side="left", padx=5)
                    ctk.CTkLabel(item, text=fecha, width=120).pack(side="left", padx=5)
                    ctk.CTkButton(item, text="Editar", width=60, height=28, fg_color="#1a73e8", command=lambda i=id_: self.editar_cliente(i)).pack(side="left", padx=2)
                    ctk.CTkButton(item, text="Eliminar", width=65, height=28, fg_color="#ff595e", command=lambda i=id_: self.eliminar_cliente(i)).pack(side="left", padx=2)
            else:
                ctk.CTkLabel(scroll_clientes, text="No hay clientes que coincidan con la búsqueda.", text_color="gray").pack(pady=30)

        txt_filtro_clientes.bind("<KeyRelease>", lambda _e: cargar_clientes_filtro(txt_filtro_clientes.get()))
        cargar_clientes_filtro()

    def editar_cliente(self, cliente_id):
        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT nombre, telefono, fecha_inscripcion FROM clientes WHERE id=?", (cliente_id,))
        data = cursor.fetchone()
        conn.close()
        if not data:
            return
        ventana = ctk.CTkToplevel(self)
        ventana.title("Editar Cliente")
        ventana.geometry("420x340")
        ventana.grab_set()
        ctk.CTkLabel(ventana, text="Nombre:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_nom = ctk.CTkEntry(ventana, width=340, font=ctk.CTkFont(size=14))
        entry_nom.insert(0, data[0])
        entry_nom.pack(pady=5)
        ctk.CTkLabel(ventana, text="Teléfono:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_tel = ctk.CTkEntry(ventana, width=340, font=ctk.CTkFont(size=14))
        entry_tel.insert(0, data[1] if data[1] else "")
        entry_tel.pack(pady=5)
        ctk.CTkLabel(ventana, text="Fecha Inscripción:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_fecha = ctk.CTkEntry(ventana, width=340, font=ctk.CTkFont(size=14))
        entry_fecha.insert(0, data[2])
        entry_fecha.pack(pady=5)

        def guardar():
            nom = entry_nom.get().strip()
            tel = entry_tel.get().strip()
            fecha = entry_fecha.get().strip()
            if not nom:
                messagebox.showerror("Error", "El nombre es obligatorio.")
                return
            if fecha:
                try:
                    datetime.strptime(fecha, "%Y-%m-%d")
                except ValueError:
                    messagebox.showerror("Error", "Fecha inválida. Use YYYY-MM-DD.")
                    return
            else:
                fecha = datetime.now().strftime("%Y-%m-%d")
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                cursor.execute("UPDATE clientes SET nombre=?, telefono=?, fecha_inscripcion=? WHERE id=?", (nom, tel, fecha, cliente_id))
                conn.commit()
                messagebox.showinfo("Éxito", "Cliente actualizado.")
                ventana.destroy()
                self.mostrar_pantalla_clientes()
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo actualizar: {e}")
            finally:
                conn.close()
        ctk.CTkButton(ventana, text="Guardar Cambios", fg_color="#a8e600", text_color="black", command=guardar).pack(pady=15)

    def eliminar_cliente(self, cliente_id):
        if messagebox.askyesno("Confirmar", "¿Eliminar este cliente? Se perderán sus datos."):
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                cursor.execute("DELETE FROM clientes WHERE id=?", (cliente_id,))
                conn.commit()
                messagebox.showinfo("Éxito", "Cliente eliminado.")
                self.mostrar_pantalla_clientes()
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo eliminar: {e}")
            finally:
                conn.close()

    # =========================================================================
    # PRODUCTOS (fila clickeable + modal de detalle con Editar/Eliminar)
    # =========================================================================
    def mostrar_pantalla_productos(self):
        self.crear_cabecera("Productos", "Catálogo maestro · Clic en una fila para ver acciones")
        main_split = ctk.CTkFrame(self.workspace, fg_color="transparent")
        main_split.pack(fill="both", expand=True)
        main_split.grid_columnconfigure(0, weight=1)
        main_split.grid_columnconfigure(1, weight=2)
        main_split.grid_rowconfigure(0, weight=1)

        form_frame = ctk.CTkFrame(main_split, fg_color="#14141e", corner_radius=12)
        form_frame.grid(row=0, column=0, padx=8, sticky="nsew")
        ctk.CTkLabel(form_frame, text="Nuevo Producto", font=ctk.CTkFont(size=16, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(20, 10))
        txt_nom = ctk.CTkEntry(form_frame, placeholder_text="Nombre del Producto", height=38, font=ctk.CTkFont(size=14))
        txt_nom.pack(fill="x", padx=20, pady=6)
        txt_cat = ctk.CTkEntry(form_frame, placeholder_text="Categoría", height=38, font=ctk.CTkFont(size=14))
        txt_cat.pack(fill="x", padx=20, pady=6)
        txt_precio = ctk.CTkEntry(form_frame, placeholder_text="Precio de Venta (RD$)", height=38, font=ctk.CTkFont(size=14))
        txt_precio.pack(fill="x", padx=20, pady=6)
        txt_costo = ctk.CTkEntry(form_frame, placeholder_text="Costo Unitario (RD$)", height=38, font=ctk.CTkFont(size=14))
        txt_costo.pack(fill="x", padx=20, pady=6)

        def acc_guardar_prod():
            nom = txt_nom.get().strip()
            cat = txt_cat.get().strip()
            prc = txt_precio.get().strip()
            cost = txt_costo.get().strip()
            if not nom or not cat or not prc or not cost:
                messagebox.showerror("Error", "Complete todos los campos.")
                return
            try:
                precio = float(prc)
                costo = float(cost)
                if precio <= 0 or costo < 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("Error", "Precio o costo inválido.")
                return
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                conn.execute("BEGIN")
                cursor.execute("INSERT INTO productos (nombre, categoria, precio_venta, costo) VALUES (?, ?, ?, ?)", (nom, cat, precio, costo))
                last_id = cursor.lastrowid
                codigo_prod = f"PROD-{last_id:04d}"
                cursor.execute("UPDATE productos SET codigo = ? WHERE id = ?", (codigo_prod, last_id))
                conn.commit()
                messagebox.showinfo("Éxito", f"Producto registrado.\nCódigo: {codigo_prod}")
                txt_nom.delete(0, "end")
                txt_cat.delete(0, "end")
                txt_precio.delete(0, "end")
                txt_costo.delete(0, "end")
            except Exception as e:
                conn.rollback()
                messagebox.showerror("Error", f"No se pudo guardar: {e}")
            finally:
                conn.close()
                self.mostrar_pantalla_productos()

        ctk.CTkButton(form_frame, text="Registrar Producto", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=14), height=40, command=acc_guardar_prod).pack(fill="x", padx=20, pady=20)

        # ---------- LISTA DE PRODUCTOS ----------
        list_frame = ctk.CTkScrollableFrame(main_split, fg_color="#14141e", corner_radius=12)
        list_frame.grid(row=0, column=1, padx=8, sticky="nsew")

        header = ctk.CTkFrame(list_frame, fg_color="#1e1e2a")
        header.pack(fill="x", padx=10, pady=(5, 2))
        ctk.CTkLabel(header, text="CÓDIGO", width=100, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="NOMBRE", width=180, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="CATEGORÍA", width=110, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="PRECIO", width=90, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="STOCK", width=70, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="VER", width=80, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, codigo, nombre, categoria, precio_venta, costo, stock FROM productos ORDER BY id DESC")
        items = cursor.fetchall()
        conn.close()

        if items:
            for id_, c, n, cat, p, cost, st in items:
                item = ctk.CTkFrame(list_frame, fg_color="#0d0d14", corner_radius=6)
                item.pack(fill="x", pady=2, padx=10)

                lbl_cod = ctk.CTkLabel(item, text=c, width=100)
                lbl_cod.pack(side="left", padx=5, pady=6)
                lbl_nom = ctk.CTkLabel(item, text=n, width=180, anchor="w")
                lbl_nom.pack(side="left", padx=5, pady=6)
                lbl_cat = ctk.CTkLabel(item, text=cat, width=110)
                lbl_cat.pack(side="left", padx=5, pady=6)
                lbl_pre = ctk.CTkLabel(item, text=f"RD${p:,.2f}", width=90)
                lbl_pre.pack(side="left", padx=5, pady=6)
                color_stock = "#ff595e" if st <= 3 else "#a8e600"
                lbl_st = ctk.CTkLabel(item, text=str(st), width=70, text_color=color_stock, font=ctk.CTkFont(weight="bold"))
                lbl_st.pack(side="left", padx=5, pady=6)

                ctk.CTkButton(
                    item, text="👁️ Ver", width=80, height=28,
                    fg_color="#333344", hover_color="#444455",
                    font=ctk.CTkFont(size=12, weight="bold"),
                    command=lambda i=id_: self.abrir_modal_detalle_producto(i)
                ).pack(side="left", padx=5, pady=6)

                for w in (item, lbl_cod, lbl_nom, lbl_cat, lbl_pre, lbl_st):
                    w.bind("<Button-1>", lambda e, i=id_: self.abrir_modal_detalle_producto(i))
        else:
            ctk.CTkLabel(list_frame, text="No hay productos registrados.", text_color="gray").pack(pady=30)

    def abrir_modal_detalle_producto(self, prod_id):
        """Muestra el detalle completo del producto y permite editarlo o eliminarlo."""
        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT codigo, nombre, categoria, precio_venta, costo, stock FROM productos WHERE id=?", (prod_id,))
        data = cursor.fetchone()
        conn.close()
        if not data:
            messagebox.showerror("Error", "Producto no encontrado.")
            return

        cod, nom, cat, precio, costo, stock = data
        ganancia = precio - costo

        modal = ctk.CTkToplevel(self)
        modal.title("Detalle de Producto")
        modal.geometry("540x560")
        modal.grab_set()
        modal.after(50, lambda: modal.focus())

        ctk.CTkLabel(modal, text="📦 DETALLE DE PRODUCTO",
                     font=ctk.CTkFont(size=18, weight="bold"),
                     text_color="#a8e600").pack(pady=(22, 5))
        ctk.CTkLabel(modal, text=nom,
                     font=ctk.CTkFont(size=22, weight="bold"),
                     text_color="white", wraplength=480).pack(pady=(0, 18))

        info = ctk.CTkFrame(modal, fg_color="#1a1a26", corner_radius=12)
        info.pack(fill="x", padx=25, pady=5)

        def agregar_fila(parent, label, valor, color="white", bold=False):
            f = ctk.CTkFrame(parent, fg_color="transparent")
            f.pack(fill="x", padx=18, pady=7)
            ctk.CTkLabel(f, text=label, font=ctk.CTkFont(size=13, weight="bold"),
                         text_color="#888899", width=140, anchor="w").pack(side="left")
            ctk.CTkLabel(f, text=valor,
                         font=ctk.CTkFont(size=14, weight="bold" if bold else "normal"),
                         text_color=color, anchor="w").pack(side="left")

        agregar_fila(info, "Código:", cod)
        agregar_fila(info, "Categoría:", cat)
        agregar_fila(info, "Precio Venta:", f"RD${precio:,.2f}", "#a8e600", bold=True)
        agregar_fila(info, "Costo:", f"RD${costo:,.2f}")
        agregar_fila(info, "Ganancia / unidad:", f"RD${ganancia:,.2f}",
                     "#a8e600" if ganancia > 0 else "#ff595e", bold=True)
        color_stock = "#ff595e" if stock <= 3 else "#a8e600"
        etiqueta_stock = f"{stock} unidades" + (" ⚠️ BAJO" if stock <= 3 else "")
        agregar_fila(info, "Stock Actual:", etiqueta_stock, color_stock, bold=True)

        ctk.CTkLabel(modal, text="¿Qué deseas hacer con este producto?",
                     font=ctk.CTkFont(size=13), text_color="gray").pack(pady=(20, 8))

        botones = ctk.CTkFrame(modal, fg_color="transparent")
        botones.pack(fill="x", padx=25, pady=5)

        def accion_editar():
            modal.destroy()
            self.editar_producto(prod_id)

        def accion_eliminar():
            modal.destroy()
            self.eliminar_producto(prod_id)

        ctk.CTkButton(botones, text="✏️ Editar", fg_color="#1a73e8", hover_color="#155bb5",
                      height=46, font=ctk.CTkFont(weight="bold", size=14),
                      command=accion_editar).pack(side="left", fill="x", expand=True, padx=(0, 5))
        ctk.CTkButton(botones, text="🗑️ Eliminar", fg_color="#ff595e", hover_color="#cc3333",
                      height=46, font=ctk.CTkFont(weight="bold", size=14),
                      command=accion_eliminar).pack(side="left", fill="x", expand=True, padx=5)

        ctk.CTkButton(modal, text="Cerrar", fg_color="#333344", hover_color="#444455",
                      height=38, font=ctk.CTkFont(weight="bold", size=13),
                      command=modal.destroy).pack(fill="x", padx=25, pady=(15, 22))

    def editar_producto(self, prod_id):
        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT nombre, categoria, precio_venta, costo, stock FROM productos WHERE id=?", (prod_id,))
        data = cursor.fetchone()
        conn.close()
        if not data:
            return

        ventana = ctk.CTkToplevel(self)
        ventana.title("Editar Producto")
        ventana.geometry("440x480")
        ventana.grab_set()
        ventana.after(50, lambda: ventana.focus())

        ctk.CTkLabel(ventana, text="Nombre:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_nom = ctk.CTkEntry(ventana, width=360, font=ctk.CTkFont(size=14))
        entry_nom.insert(0, data[0])
        entry_nom.pack(pady=5)

        ctk.CTkLabel(ventana, text="Categoría:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_cat = ctk.CTkEntry(ventana, width=360, font=ctk.CTkFont(size=14))
        entry_cat.insert(0, data[1])
        entry_cat.pack(pady=5)

        ctk.CTkLabel(ventana, text="Precio Venta (RD$):", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_precio = ctk.CTkEntry(ventana, width=360, font=ctk.CTkFont(size=14))
        entry_precio.insert(0, str(data[2]))
        entry_precio.pack(pady=5)

        ctk.CTkLabel(ventana, text="Costo (RD$):", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_costo = ctk.CTkEntry(ventana, width=360, font=ctk.CTkFont(size=14))
        entry_costo.insert(0, str(data[3]))
        entry_costo.pack(pady=5)

        ctk.CTkLabel(ventana, text="Stock:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_stock = ctk.CTkEntry(ventana, width=360, font=ctk.CTkFont(size=14))
        entry_stock.insert(0, str(data[4]))
        entry_stock.pack(pady=5)

        def guardar():
            nom = entry_nom.get().strip()
            cat = entry_cat.get().strip()
            try:
                precio = float(entry_precio.get().strip())
                costo = float(entry_costo.get().strip())
                stock = int(entry_stock.get().strip())
                if precio <= 0 or costo < 0 or stock < 0:
                    raise ValueError
            except Exception:
                messagebox.showerror("Error", "Precio, costo y stock deben ser números válidos.")
                return
            if not nom or not cat:
                messagebox.showerror("Error", "Nombre y categoría son obligatorios.")
                return
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                cursor.execute("UPDATE productos SET nombre=?, categoria=?, precio_venta=?, costo=?, stock=? WHERE id=?",
                               (nom, cat, precio, costo, stock, prod_id))
                conn.commit()
                messagebox.showinfo("Éxito", "Producto actualizado.")
                ventana.destroy()
                self.mostrar_pantalla_productos()
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo actualizar: {e}")
            finally:
                conn.close()

        ctk.CTkButton(ventana, text="Guardar Cambios", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold"), command=guardar).pack(pady=15)

    def eliminar_producto(self, prod_id):
        if messagebox.askyesno("Confirmar", "¿Eliminar este producto? Se perderán todos los datos."):
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                cursor.execute("DELETE FROM productos WHERE id=?", (prod_id,))
                conn.commit()
                messagebox.showinfo("Éxito", "Producto eliminado.")
                self.mostrar_pantalla_productos()
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo eliminar: {e}")
            finally:
                conn.close()

    # =========================================================================
    # INVENTARIO
    # =========================================================================
    def mostrar_pantalla_inventario(self):
        self.crear_cabecera("Inventario", "Registro de compras y stock actual")
        main_split = ctk.CTkFrame(self.workspace, fg_color="transparent")
        main_split.pack(fill="both", expand=True)
        main_split.grid_columnconfigure(0, weight=1)
        main_split.grid_columnconfigure(1, weight=1)
        main_split.grid_rowconfigure(0, weight=1)

        form_frame = ctk.CTkFrame(main_split, fg_color="#14141e", corner_radius=12)
        form_frame.grid(row=0, column=0, padx=8, sticky="nsew")
        ctk.CTkLabel(form_frame, text="Entrada de Mercancía", font=ctk.CTkFont(size=16, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(20, 10))

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT codigo, nombre FROM productos")
        prods_db = cursor.fetchall()
        conn.close()
        lista_opciones = [f"{p[0]} - {p[1]}" for p in prods_db] if prods_db else ["No hay productos"]

        ctk.CTkLabel(form_frame, text="Producto:", text_color="gray").pack(anchor="w", padx=20, pady=(5, 2))
        cmb_prod_inv = ctk.CTkOptionMenu(form_frame, values=lista_opciones, height=38, font=ctk.CTkFont(size=14))
        cmb_prod_inv.pack(fill="x", padx=20, pady=6)
        txt_cant_inv = ctk.CTkEntry(form_frame, placeholder_text="Cantidad Entrante", height=38, font=ctk.CTkFont(size=14))
        txt_cant_inv.pack(fill="x", padx=20, pady=6)
        txt_costo_inv = ctk.CTkEntry(form_frame, placeholder_text="Costo Unitario (RD$)", height=38, font=ctk.CTkFont(size=14))
        txt_costo_inv.pack(fill="x", padx=20, pady=6)

        def acc_grabar_compra():
            seleccion = cmb_prod_inv.get()
            cant_str = txt_cant_inv.get().strip()
            costo_str = txt_costo_inv.get().strip()
            if seleccion == "No hay productos":
                messagebox.showerror("Error", "Cree productos primero.")
                return
            if not cant_str or not costo_str:
                messagebox.showerror("Error", "Complete todos los campos.")
                return
            try:
                cant = int(cant_str)
                costo = float(costo_str)
                if cant <= 0 or costo < 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("Error", "Cantidad y costo deben ser válidos.")
                return
            cod_prod, nom_prod = seleccion.split(" - ")
            fecha_ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                conn.execute("BEGIN")
                cursor.execute("SELECT id FROM productos WHERE codigo=?", (cod_prod,))
                id_prod = cursor.fetchone()[0]
                cursor.execute("INSERT INTO compras (producto_info, cantidad, costo_producto, fecha_compra) VALUES (?, ?, ?, ?)", (seleccion, cant, costo, fecha_ahora))
                cursor.execute("UPDATE productos SET stock = stock + ? WHERE codigo = ?", (cant, cod_prod))
                cursor.execute("INSERT INTO historial_inventario (producto_id, nombre_producto, tipo_movimiento, cantidad, fecha, usuario, motivo) VALUES (?, ?, 'ENTRADA', ?, ?, ?, ?)", (id_prod, nom_prod, cant, fecha_ahora, self.usuario_actual, "Compra/Ingreso a stock"))
                conn.commit()
                messagebox.showinfo("Éxito", "Compra registrada y stock actualizado.")
                txt_cant_inv.delete(0, "end")
                txt_costo_inv.delete(0, "end")
            except Exception as e:
                conn.rollback()
                messagebox.showerror("Error", f"Falló la compra: {e}")
            finally:
                conn.close()
                self.mostrar_pantalla_inventario()

        ctk.CTkButton(form_frame, text="Grabar Entrada de Compra", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=14), height=40, command=acc_grabar_compra).pack(fill="x", padx=20, pady=20)

        right_frame = ctk.CTkFrame(main_split, fg_color="#14141e", corner_radius=12)
        right_frame.grid(row=0, column=1, padx=8, sticky="nsew")
        ctk.CTkLabel(right_frame, text="Stock Actual", font=ctk.CTkFont(size=16, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(15, 5))
        scroll_inv = ctk.CTkScrollableFrame(right_frame, fg_color="#1a1a26")
        scroll_inv.pack(fill="both", expand=True, padx=15, pady=15)

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT codigo, nombre, stock FROM productos")
        stocks = cursor.fetchall()
        conn.close()
        if stocks:
            for cp, np, st in stocks:
                f = ctk.CTkFrame(scroll_inv, fg_color="#0d0d14")
                f.pack(fill="x", pady=2, padx=5)
                ctk.CTkLabel(f, text=f"{cp} - {np}", font=ctk.CTkFont(size=14)).pack(side="left", padx=10)
                ctk.CTkLabel(f, text=f"Stock: {st} uds", text_color="#a8e600", font=ctk.CTkFont(weight="bold", size=14)).pack(side="right", padx=15)
        else:
            ctk.CTkLabel(scroll_inv, text="No hay productos en el inventario.", text_color="gray").pack(pady=20)

        ctk.CTkLabel(right_frame, text="Historial de Compras", font=ctk.CTkFont(size=16, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(15, 5))
        scroll_compras = ctk.CTkScrollableFrame(right_frame, fg_color="#1a1a26", height=200)
        scroll_compras.pack(fill="both", expand=True, padx=15, pady=15)

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, producto_info, cantidad, costo_producto, fecha_compra FROM compras ORDER BY id DESC")
        compras = cursor.fetchall()
        conn.close()

        if compras:
            for c_id, info, cant, costo, fecha in compras:
                f = ctk.CTkFrame(scroll_compras, fg_color="#0d0d14")
                f.pack(fill="x", pady=2, padx=5)
                ctk.CTkLabel(f, text=f"{info} | {cant} uds | RD${costo:,.2f} c/u | {fecha}", font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
                ctk.CTkButton(f, text="Editar", width=60, height=26, fg_color="#1a73e8", command=lambda i=c_id: self.editar_compra(i)).pack(side="right", padx=2)
                ctk.CTkButton(f, text="Eliminar", width=60, height=26, fg_color="#ff595e", command=lambda i=c_id: self.eliminar_compra(i)).pack(side="right", padx=2)
        else:
            ctk.CTkLabel(scroll_compras, text="No hay compras registradas.", text_color="gray").pack(pady=20)

    def editar_compra(self, compra_id):
        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT producto_info, cantidad, costo_producto FROM compras WHERE id=?", (compra_id,))
        data = cursor.fetchone()
        conn.close()
        if not data:
            return
        ventana = ctk.CTkToplevel(self)
        ventana.title("Editar Compra")
        ventana.geometry("420x280")
        ventana.grab_set()
        ctk.CTkLabel(ventana, text="Producto:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_info = ctk.CTkEntry(ventana, width=340, font=ctk.CTkFont(size=14))
        entry_info.insert(0, data[0])
        entry_info.pack(pady=5)
        ctk.CTkLabel(ventana, text="Cantidad:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_cant = ctk.CTkEntry(ventana, width=340, font=ctk.CTkFont(size=14))
        entry_cant.insert(0, str(data[1]))
        entry_cant.pack(pady=5)
        ctk.CTkLabel(ventana, text="Costo Unitario (RD$):", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_costo = ctk.CTkEntry(ventana, width=340, font=ctk.CTkFont(size=14))
        entry_costo.insert(0, str(data[2]))
        entry_costo.pack(pady=5)

        def guardar():
            try:
                cant = int(entry_cant.get().strip())
                costo = float(entry_costo.get().strip())
                if cant <= 0 or costo < 0:
                    raise ValueError
            except:
                messagebox.showerror("Error", "Cantidad y costo deben ser números válidos.")
                return
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                cursor.execute("UPDATE compras SET producto_info=?, cantidad=?, costo_producto=? WHERE id=?", (entry_info.get().strip(), cant, costo, compra_id))
                conn.commit()
                messagebox.showinfo("Éxito", "Compra actualizada.")
                ventana.destroy()
                self.mostrar_pantalla_inventario()
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo actualizar: {e}")
            finally:
                conn.close()
        ctk.CTkButton(ventana, text="Guardar Cambios", fg_color="#a8e600", text_color="black", command=guardar).pack(pady=15)

    def eliminar_compra(self, compra_id):
        if messagebox.askyesno("Confirmar", "¿Eliminar esta compra? El stock se ajustará automáticamente."):
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                conn.execute("BEGIN")
                cursor.execute("SELECT producto_info, cantidad FROM compras WHERE id=?", (compra_id,))
                info, cant = cursor.fetchone()
                cod_prod, nom_prod = info.split(" - ")
                cursor.execute("SELECT id FROM productos WHERE codigo=?", (cod_prod,))
                id_prod = cursor.fetchone()[0]
                cursor.execute("UPDATE productos SET stock = stock - ? WHERE codigo = ?", (cant, cod_prod))
                cursor.execute("DELETE FROM compras WHERE id=?", (compra_id,))
                fecha_ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                cursor.execute("INSERT INTO historial_inventario (producto_id, nombre_producto, tipo_movimiento, cantidad, fecha, usuario, motivo) VALUES (?, ?, 'SALIDA', ?, ?, ?, ?)", (id_prod, nom_prod, cant, fecha_ahora, self.usuario_actual, "Eliminación de registro de compra"))
                conn.commit()
                messagebox.showinfo("Éxito", "Compra eliminada y stock ajustado.")
                self.mostrar_pantalla_inventario()
            except Exception as e:
                conn.rollback()
                messagebox.showerror("Error", f"No se pudo eliminar: {e}")
            finally:
                conn.close()

    # =========================================================================
    # CONTROL DE CAJAS
    # =========================================================================
    def mostrar_pantalla_control_caja(self):
        self.crear_cabecera("Control de Caja", "Apertura y cierre de turnos manual")
        frame_box = ctk.CTkFrame(self.workspace, fg_color="#14141e", corner_radius=14)
        frame_box.pack(fill="both", expand=True, padx=10, pady=10)

        if not self.id_caja_activa:
            ctk.CTkLabel(frame_box, text="PUNTO DE VENTA CERRADO", font=ctk.CTkFont(size=20, weight="bold"), text_color="#ff595e").pack(pady=(40, 10))
            ctk.CTkLabel(frame_box, text="Ingrese el monto con el que inicia su turno en caja:", font=ctk.CTkFont(size=14), text_color="gray").pack(pady=10)
            txt_monto_inicial = ctk.CTkEntry(frame_box, placeholder_text="Ej: 500.0", width=200, height=40, font=ctk.CTkFont(size=15))
            txt_monto_inicial.pack(pady=15)

            def acc_abrir_caja():
                monto_str = txt_monto_inicial.get().strip()
                if not monto_str:
                    monto_apertura = 0.0
                else:
                    try:
                        monto_apertura = float(monto_str)
                    except ValueError:
                        messagebox.showerror("Error", "El monto ingresado no es válido.")
                        return
                ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                conn = self.conectar_db()
                cursor = conn.cursor()
                cursor.execute("INSERT INTO cajas (usuario, monto_apertura, fecha_apertura, estado) VALUES (?, ?, ?, 'Abierta')", (self.usuario_actual, monto_apertura, ahora))
                self.id_caja_activa = cursor.lastrowid
                conn.commit()
                conn.close()
                self._actualizar_indicador_caja()  # <-- Actualizar indicador
                messagebox.showinfo("Éxito", f"Caja aperturada correctamente con RD${monto_apertura:,.2f}.")
                self.mostrar_pantalla_control_caja()

            ctk.CTkButton(frame_box, text="ABRIR TURNO", fg_color="#a8e600", text_color="black", height=48, font=ctk.CTkFont(weight="bold", size=15), command=acc_abrir_caja).pack(pady=20)
        else:
            ctk.CTkLabel(frame_box, text=f"TURNO ABIERTO (ID: {self.id_caja_activa})", font=ctk.CTkFont(size=18, weight="bold"), text_color="#a8e600").pack(pady=(40, 10))
            conn = self.conectar_db()
            cursor = conn.cursor()
            cursor.execute("SELECT monto_apertura FROM cajas WHERE id=?", (self.id_caja_activa,))
            monto = cursor.fetchone()[0]
            conn.close()
            ctk.CTkLabel(frame_box, text=f"Monto Inicial Registrado: RD${monto:,.2f}", font=ctk.CTkFont(size=14), text_color="gray").pack(pady=10)
            ctk.CTkButton(frame_box, text="HACER CUADRE Y CERRAR TURNO", fg_color="#ff595e", text_color="white", height=48, font=ctk.CTkFont(weight="bold", size=15), command=self.mostrar_pantalla_corte).pack(pady=20)

    # =========================================================================
    # CORTE DE CAJA
    # =========================================================================
    def mostrar_pantalla_corte(self):
        self.crear_cabecera("Corte de Caja", "Resumen detallado del turno actual")
        if not self.id_caja_activa:
            ctk.CTkLabel(self.workspace, text="No hay caja abierta.", text_color="#ff595e").pack(pady=60)
            return

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT monto_apertura, fecha_apertura FROM cajas WHERE id=?", (self.id_caja_activa,))
        monto_ap, fecha_ap = cursor.fetchone()
        fecha_ap_dt = datetime.strptime(fecha_ap, "%Y-%m-%d %H:%M:%S")

        cursor.execute("SELECT metodo_pago, SUM(total) FROM facturas WHERE caja_id=? GROUP BY metodo_pago", (self.id_caja_activa,))
        pagos = cursor.fetchall()
        total_efectivo = total_transferencia = 0.0
        for metodo, monto in pagos:
            if metodo == "Efectivo": total_efectivo += monto
            elif metodo == "Transferencia": total_transferencia += monto

        cursor.execute("""
            SELECT p.categoria, COALESCE(SUM(fd.cantidad * fd.precio_unitario), 0)
            FROM factura_detalles fd JOIN facturas f ON fd.factura_id = f.id
            JOIN productos p ON fd.nombre_producto = p.nombre
            WHERE f.caja_id=? GROUP BY p.categoria
        """, (self.id_caja_activa,))
        ventas_depto = cursor.fetchall()

        cursor.execute("""
            SELECT fd.nombre_producto, SUM(fd.cantidad), SUM(fd.cantidad * fd.precio_unitario)
            FROM factura_detalles fd JOIN facturas f ON fd.factura_id = f.id
            WHERE f.caja_id=? GROUP BY fd.nombre_producto ORDER BY fd.nombre_producto
        """, (self.id_caja_activa,))
        detalle_productos = cursor.fetchall()

        cursor.execute("""
            SELECT COALESCE(SUM(fd.cantidad * (fd.precio_unitario - p.costo)), 0)
            FROM factura_detalles fd JOIN facturas f ON fd.factura_id = f.id
            JOIN productos p ON fd.nombre_producto = p.nombre WHERE f.caja_id=?
        """, (self.id_caja_activa,))
        ganancia = cursor.fetchone()[0] or 0.0
        total_ventas = sum(monto for _, monto in pagos) if pagos else 0.0
        cursor.execute("SELECT COUNT(*) FROM facturas WHERE caja_id=?", (self.id_caja_activa,))
        num_ventas = cursor.fetchone()[0]
        total_efectivo_caja = monto_ap + total_efectivo
        conn.close()

        main_frame = ctk.CTkFrame(self.workspace, fg_color="transparent")
        main_frame.pack(fill="both", expand=True)
        scroll_frame = ctk.CTkScrollableFrame(main_frame, fg_color="#14141e", corner_radius=14)
        scroll_frame.pack(fill="both", expand=True, padx=5, pady=5)

        ctk.CTkLabel(scroll_frame, text="MONSTER GYM - CORTE DE CAJA", font=ctk.CTkFont(size=22, weight="bold"), text_color="#a8e600").pack(pady=12)
        ctk.CTkLabel(scroll_frame, text=f"Fecha: {fecha_ap_dt.strftime('%d/%m/%Y %H:%M')}", font=ctk.CTkFont(size=15)).pack()
        ctk.CTkLabel(scroll_frame, text=f"Cajero: {self.usuario_actual}", font=ctk.CTkFont(size=15)).pack()
        ctk.CTkLabel(scroll_frame, text=f"Caja ID: {self.id_caja_activa}", font=ctk.CTkFont(size=15)).pack(pady=(0, 12))
        ctk.CTkLabel(scroll_frame, text=f"Fondo Base (Apertura): RD${monto_ap:,.2f}", font=ctk.CTkFont(size=15, weight="bold")).pack(anchor="w", padx=20, pady=5)
        ctk.CTkLabel(scroll_frame, text="--- INGRESOS REGISTRADOS ---", font=ctk.CTkFont(size=15, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(10, 2))
        ctk.CTkLabel(scroll_frame, text=f"Ventas en Efectivo: RD${total_efectivo:,.2f}").pack(anchor="w", padx=30)
        ctk.CTkLabel(scroll_frame, text=f"Transferencia: RD${total_transferencia:,.2f}").pack(anchor="w", padx=30)
        if ventas_depto:
            ctk.CTkLabel(scroll_frame, text="--- VENTAS POR DEPARTAMENTO ---", font=ctk.CTkFont(size=15, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(10, 2))
            for cat, total in ventas_depto:
                ctk.CTkLabel(scroll_frame, text=f"{cat}: RD${total:,.2f}").pack(anchor="w", padx=30)
        if detalle_productos:
            ctk.CTkLabel(scroll_frame, text="--- DETALLE DE PRODUCTOS VENDIDOS ---", font=ctk.CTkFont(size=15, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(10, 2))
            for prod, cant, subt in detalle_productos:
                ctk.CTkLabel(scroll_frame, text=f"{prod} x{cant} = RD${subt:,.2f}").pack(anchor="w", padx=30)
        ctk.CTkLabel(scroll_frame, text="--- TOTALES ---", font=ctk.CTkFont(size=15, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(10, 2))
        ctk.CTkLabel(scroll_frame, text=f"Ventas Totales: RD${total_ventas:,.2f}", font=ctk.CTkFont(size=15, weight="bold")).pack(anchor="w", padx=30)
        ctk.CTkLabel(scroll_frame, text=f"Efectivo Esperado en Caja: RD${total_efectivo_caja:,.2f}", font=ctk.CTkFont(size=16, weight="bold"), text_color="#ffca3a").pack(anchor="w", padx=30, pady=5)
        ctk.CTkLabel(scroll_frame, text="--- FIRMAS ---", font=ctk.CTkFont(size=15, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(10, 2))
        ctk.CTkLabel(scroll_frame, text="Cajero: _________________________").pack(anchor="w", padx=30)
        ctk.CTkLabel(scroll_frame, text="Recibe: _________________________").pack(anchor="w", padx=30)

        def imprimir_corte_y_cerrar():
            ANCHO = 48
            def fmt(texto, alineacion="left"):
                if len(texto) >= ANCHO:
                    return texto[:ANCHO]
                if alineacion == "center":
                    return texto.center(ANCHO)
                elif alineacion == "right":
                    return texto.rjust(ANCHO)
                else:
                    return texto.ljust(ANCHO)

            ticket = ""
            ticket += fmt("=" * ANCHO) + "\n"
            ticket += fmt("MONSTER GYM", "center") + "\n"
            ticket += fmt("C/PRINCIPAL NUEVA BARQUITA", "center") + "\n"
            ticket += fmt("SEGUNDO NIVEL", "center") + "\n"
            ticket += fmt("RFC0031282AB1", "center") + "\n"
            ticket += fmt("CONTACTO 849 228 7784", "center") + "\n"
            ticket += fmt("CORTE DEL DIA", "center") + "\n"
            ticket += fmt(f"DEL {fecha_ap_dt.strftime('%d/%b/%Y').upper()}", "center") + "\n"
            ticket += fmt("-" * ANCHO) + "\n"
            ticket += fmt(f"REALIZADO: {fecha_ap_dt.strftime('%d/%b/%Y %I:%M %p').upper()}", "left") + "\n"
            ticket += fmt(f"CAJERO: {self.usuario_actual.upper()}", "left") + "\n"
            ticket += fmt(f"VENTAS: {num_ventas}", "left") + "\n"
            ticket += fmt("=" * ANCHO) + "\n"
            ticket += fmt(f"INICIO CAJA: $RD{monto_ap:,.2f}", "left") + "\n"
            ticket += fmt(f"EFECTIVO VENTAS: $RD{total_efectivo:,.2f}", "left") + "\n"
            ticket += fmt(f"TRANSFERENCIA: $RD{total_transferencia:,.2f}", "left") + "\n"
            ticket += fmt(f"TOTAL CONTADO: $RD{total_efectivo + total_transferencia:,.2f}", "left") + "\n"
            ticket += fmt("=" * ANCHO) + "\n"
            ticket += fmt(f"EFECTIVO EN CAJA: $RD{monto_ap + total_efectivo:,.2f}", "left") + "\n"
            ticket += fmt("=" * ANCHO) + "\n"
            ticket += fmt(f"VENTAS TOTALES: $RD{total_ventas:,.2f}", "left") + "\n"
            ticket += fmt(f"GANANCIA: $RD{ganancia:,.2f}", "left") + "\n"
            ticket += fmt("=" * ANCHO) + "\n"
            if detalle_productos:
                ticket += fmt("PRODUCTOS VENDIDOS", "center") + "\n"
                for prod, cant, subt in detalle_productos:
                    ticket += fmt(f"{prod} x{cant} = $RD{subt:,.2f}", "left") + "\n"
                ticket += fmt("=" * ANCHO) + "\n"
            else:
                ticket += fmt("SIN PRODUCTOS VENDIDOS", "center") + "\n"
            ticket += fmt("Cajero: _________________", "left") + "\n"
            ticket += fmt("Encargado: _________________", "left") + "\n"
            ticket += fmt("=" * ANCHO) + "\n"

            if self.impresora:
                try:
                    self.impresora.imprimir(ticket)
                    messagebox.showinfo("Impresión", "Ticket de corte enviado a la impresora.")
                except Exception as e:
                    messagebox.showerror("Error de impresión", f"No se pudo imprimir:\n{e}")
                    print(ticket)
            else:
                print(ticket)
                messagebox.showinfo("Impresión (simulada)", "Ticket impreso en consola.")

            ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            conn = self.conectar_db()
            cursor = conn.cursor()
            cursor.execute("UPDATE cajas SET monto_cierre=?, fecha_cierre=?, estado='Cerrada' WHERE id=?", (total_ventas, ahora, self.id_caja_activa))
            conn.commit()
            conn.close()
            self.id_caja_activa = None
            self._actualizar_indicador_caja()  # <-- Actualizar indicador
            messagebox.showinfo("Cierre exitoso", "El turno se ha cerrado correctamente.")
            self.mostrar_dashboard()

        ctk.CTkButton(scroll_frame, text="Imprimir Corte y Cerrar Turno", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=15), height=44, command=imprimir_corte_y_cerrar).pack(pady=20)

    # =========================================================================
    # FACTURACIÓN
    # =========================================================================
    def mostrar_pantalla_facturacion(self):
        self.crear_cabecera("Facturación", "Venta actual")
        self.carrito_actual = []

        if not self.id_caja_activa:
            ctk.CTkLabel(self.workspace, text="CAJA CERRADA: Debe abrir un turno para facturar.", font=ctk.CTkFont(size=16, weight="bold"), text_color="#ff595e").pack(pady=60)
            return

        main_frame = ctk.CTkFrame(self.workspace, fg_color="transparent")
        main_frame.pack(fill="both", expand=True)
        main_frame.grid_columnconfigure(0, weight=3)
        main_frame.grid_columnconfigure(1, weight=1)
        main_frame.grid_rowconfigure(0, weight=1)

        # -------- IZQUIERDA --------
        left_panel = ctk.CTkFrame(main_frame, fg_color="#14141e", corner_radius=12)
        left_panel.grid(row=0, column=0, padx=8, sticky="nsew")
        left_panel.grid_rowconfigure(0, weight=0)
        left_panel.grid_rowconfigure(1, weight=1)

        selector_frame = ctk.CTkFrame(left_panel, fg_color="transparent")
        selector_frame.grid(row=0, column=0, pady=10, padx=10, sticky="ew")

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, codigo, nombre, precio_venta, stock FROM productos ORDER BY nombre")
        prods_db = cursor.fetchall()
        conn.close()

        mapeo_productos = {f"{p[1]} - {p[2]} (RD${p[3]:,.2f})": p[0] for p in prods_db if p[4] > 0}
        lista_nombres_prod = list(mapeo_productos.keys()) if mapeo_productos else ["No hay productos con stock"]

        ctk.CTkLabel(selector_frame, text="Agregar producto:", font=ctk.CTkFont(size=13), text_color="gray").grid(row=0, column=0, sticky="w", padx=(0, 10))
        cmb_prod = ctk.CTkOptionMenu(selector_frame, values=lista_nombres_prod, height=36, width=330, font=ctk.CTkFont(size=13))
        cmb_prod.grid(row=0, column=1, padx=5)
        txt_cant = ctk.CTkEntry(selector_frame, placeholder_text="Cant.", width=80, height=36, font=ctk.CTkFont(size=13))
        txt_cant.insert(0, "1")
        txt_cant.grid(row=0, column=2, padx=5)

        carrito_frame = ctk.CTkFrame(left_panel, fg_color="#1a1a26", corner_radius=10)
        carrito_frame.grid(row=1, column=0, padx=10, pady=(0, 10), sticky="nsew")
        carrito_frame.grid_columnconfigure(0, weight=1)
        carrito_frame.grid_rowconfigure(1, weight=1)

        header_carrito = ctk.CTkFrame(carrito_frame, fg_color="transparent")
        header_carrito.grid(row=0, column=0, sticky="ew", padx=10, pady=(8, 2))
        ctk.CTkLabel(header_carrito, text="Producto", width=250, anchor="w", font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header_carrito, text="Cant.", width=60, anchor="w", font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header_carrito, text="Acción", width=80, anchor="center", font=ctk.CTkFont(weight="bold")).pack(side="right", padx=15)
        ctk.CTkLabel(header_carrito, text="Total", width=90, anchor="e", font=ctk.CTkFont(weight="bold")).pack(side="right", padx=5)

        items_scroll = ctk.CTkScrollableFrame(carrito_frame, fg_color="transparent", height=250)
        items_scroll.grid(row=1, column=0, sticky="nsew", padx=5, pady=5)
        items_scroll.grid_columnconfigure(0, weight=1)

        def actualizar_totales():
            total = sum(item["precio"] * item["cantidad"] for item in self.carrito_actual)
            lbl_total.configure(text=f"RD${total:,.2f}")
            calcular_cambio()

        def calcular_cambio():
            total = sum(item["precio"] * item["cantidad"] for item in self.carrito_actual)
            try:
                recibido = float(txt_recibido.get().strip())
            except (TypeError, ValueError):
                recibido = 0.0
            cambio = max(0, recibido - total)
            lbl_cambio.configure(text=f"RD${cambio:,.2f}")

        def eliminar_item_carrito(index):
            if 0 <= index < len(self.carrito_actual):
                del self.carrito_actual[index]
                refrescar_tabla_carrito()

        def refrescar_tabla_carrito():
            for widget in items_scroll.winfo_children():
                widget.destroy()
            if not self.carrito_actual:
                ctk.CTkLabel(items_scroll, text="Sin productos", text_color="gray", font=ctk.CTkFont(size=14)).grid(row=0, column=0, pady=30)
            else:
                for idx, item in enumerate(self.carrito_actual):
                    frame_item = ctk.CTkFrame(items_scroll, fg_color="#0d0d14", corner_radius=6)
                    frame_item.grid(row=idx, column=0, sticky="ew", pady=2, padx=2)
                    frame_item.grid_columnconfigure(0, weight=1)
                    subt = item["precio"] * item["cantidad"]
                    ctk.CTkLabel(frame_item, text=item["nombre"], width=250, anchor="w", font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
                    ctk.CTkLabel(frame_item, text=str(item["cantidad"]), width=60, anchor="w", font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
                    ctk.CTkButton(frame_item, text="✕ Eliminar", width=85, height=28, fg_color="#ff595e", hover_color="#cc3333",
                                  font=ctk.CTkFont(size=11, weight="bold"),
                                  command=lambda i=idx: eliminar_item_carrito(i)).pack(side="right", padx=8)
                    ctk.CTkLabel(frame_item, text=f"RD${subt:,.2f}", width=90, anchor="e", font=ctk.CTkFont(size=13)).pack(side="right", padx=5)
            actualizar_totales()

        def acc_agregar_producto():
            sel = cmb_prod.get()
            qty_str = txt_cant.get().strip()
            if sel not in mapeo_productos:
                messagebox.showerror("Error", "Seleccione un producto válido.")
                return
            try:
                qty = int(qty_str)
                if qty <= 0:
                    raise ValueError
            except (TypeError, ValueError):
                messagebox.showerror("Error", "Cantidad inválida.")
                return
            prod_id = mapeo_productos[sel]
            conn = self.conectar_db()
            cursor = conn.cursor()
            cursor.execute("SELECT codigo, nombre, precio_venta, stock FROM productos WHERE id=?", (prod_id,))
            prod = cursor.fetchone()
            conn.close()
            if not prod:
                messagebox.showerror("Error", "Producto no encontrado.")
                return
            cantidad_actual = sum(i["cantidad"] for i in self.carrito_actual if i["id"] == prod_id)
            disponible = prod[3] - cantidad_actual
            if disponible < qty:
                messagebox.showerror("Stock insuficiente", f"Disponible: {disponible} unidades.")
                return
            for item in self.carrito_actual:
                if item["id"] == prod_id:
                    item["cantidad"] += qty
                    break
            else:
                self.carrito_actual.append({"id": prod_id, "codigo": prod[0], "nombre": prod[1], "precio": prod[2], "cantidad": qty})
            refrescar_tabla_carrito()
            txt_cant.delete(0, "end")
            txt_cant.insert(0, "1")

        btn_add = ctk.CTkButton(selector_frame, text="＋ Añadir", width=100, height=36, fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=13), command=acc_agregar_producto)
        btn_add.grid(row=0, column=3, padx=5)
        txt_cant.bind("<Return>", lambda _e: acc_agregar_producto())

        # -------- DERECHA --------
        right_panel = ctk.CTkFrame(main_frame, fg_color="#14141e", corner_radius=12)
        right_panel.grid(row=0, column=1, padx=8, sticky="nsew")

        ctk.CTkLabel(right_panel, text="Cliente", font=ctk.CTkFont(size=13), text_color="gray").pack(anchor="w", padx=20, pady=(20, 2))

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, codigo, nombre, fecha_inscripcion, fecha_vencimiento FROM clientes ORDER BY nombre")
        clientes_db = cursor.fetchall()
        conn.close()

        clientes_map = {f"{c[1]} - {c[2]}": c for c in clientes_db}
        cliente_seleccionado = {"valor": "CLI-0000 - Consumidor Final"}

        txt_buscar_cliente = ctk.CTkEntry(right_panel, placeholder_text="Buscar por nombre o código...", height=36, font=ctk.CTkFont(size=13))
        txt_buscar_cliente.pack(fill="x", padx=20, pady=(0, 5))

        cliente_resultados = ctk.CTkScrollableFrame(right_panel, fg_color="#0d0d14", height=90, corner_radius=8)

        lbl_cliente_actual = ctk.CTkLabel(right_panel, text="Cliente: Consumidor Final", font=ctk.CTkFont(size=13, weight="bold"), text_color="#a8e600")
        lbl_cliente_actual.pack(anchor="w", padx=20, pady=(0, 4))

        lbl_estado_cliente = ctk.CTkLabel(right_panel, text="", font=ctk.CTkFont(size=12), text_color="gray")
        lbl_estado_cliente.pack(anchor="w", padx=20, pady=(0, 8))

        def seleccionar_cliente(valor_cliente):
            cliente_seleccionado["valor"] = valor_cliente
            txt_buscar_cliente.delete(0, "end")
            txt_buscar_cliente.insert(0, valor_cliente)
            cliente_resultados.pack_forget()
            for widget in cliente_resultados.winfo_children():
                widget.destroy()
            datos = clientes_map.get(valor_cliente)
            if datos:
                _id, codigo, nombre, fecha_insc, fecha_vence = datos
                lbl_cliente_actual.configure(text=f"Cliente: {nombre}")
                if fecha_vence:
                    hoy = datetime.now().date()
                    vence = datetime.strptime(fecha_vence, "%Y-%m-%d").date()
                    dias = (hoy - vence).days
                    if dias > 0:
                        lbl_estado_cliente.configure(text=f"Suscripción vencida: {dias} día(s)", text_color="#ff595e")
                    else:
                        lbl_estado_cliente.configure(text=f"Vence: {fecha_vence}", text_color="#a8e600")
                else:
                    lbl_estado_cliente.configure(text=f"Inscrito: {fecha_insc} · Sin suscripción", text_color="gray")

        def filtrar_clientes_facturacion(_event=None):
            texto = txt_buscar_cliente.get().strip().lower()
            if not texto or texto == cliente_seleccionado["valor"].lower():
                cliente_resultados.pack_forget()
                for widget in cliente_resultados.winfo_children():
                    widget.destroy()
                return

            cliente_resultados.pack(fill="x", padx=20, pady=(0, 8))
            for widget in cliente_resultados.winfo_children():
                widget.destroy()

            coincidencias = [(clave, datos) for clave, datos in clientes_map.items() if texto in clave.lower()]
            if not coincidencias:
                ctk.CTkLabel(cliente_resultados, text="No se encontraron clientes.", text_color="gray").pack(pady=10)
                return
            for clave, datos in coincidencias[:30]:
                ctk.CTkButton(cliente_resultados, text=f"{datos[1]} - {datos[2]}", anchor="w", height=28, fg_color="transparent",
                              hover_color="#1f1f2a", font=ctk.CTkFont(size=12),
                              command=lambda v=clave: seleccionar_cliente(v)).pack(fill="x", padx=4, pady=1)

        txt_buscar_cliente.bind("<KeyRelease>", filtrar_clientes_facturacion)
        seleccionar_cliente("CLI-0000 - Consumidor Final")

        ctk.CTkLabel(right_panel, text="Método de pago", font=ctk.CTkFont(size=13), text_color="gray").pack(anchor="w", padx=20, pady=(10, 2))
        cmb_metodo_pago = ctk.CTkOptionMenu(right_panel, values=["Efectivo", "Transferencia"], height=36, width=200, font=ctk.CTkFont(size=13))
        cmb_metodo_pago.pack(fill="x", padx=20, pady=2)
        cmb_metodo_pago.set("Efectivo")

        ctk.CTkLabel(right_panel, text="Recibido", font=ctk.CTkFont(size=13), text_color="gray").pack(anchor="w", padx=20, pady=(15, 2))
        txt_recibido = ctk.CTkEntry(right_panel, placeholder_text="0", height=40, width=200, font=ctk.CTkFont(size=14))
        txt_recibido.pack(fill="x", padx=20, pady=2)
        txt_recibido.insert(0, "0")

        ctk.CTkLabel(right_panel, text="Cambio", font=ctk.CTkFont(size=13), text_color="gray").pack(anchor="w", padx=20, pady=(10, 2))
        lbl_cambio = ctk.CTkLabel(right_panel, text="RD$0.00", font=ctk.CTkFont(size=18, weight="bold"), text_color="white")
        lbl_cambio.pack(anchor="w", padx=20)

        ctk.CTkLabel(right_panel, text="Total", font=ctk.CTkFont(size=13), text_color="gray").pack(anchor="w", padx=20, pady=(15, 2))
        lbl_total = ctk.CTkLabel(right_panel, text="RD$0.00", font=ctk.CTkFont(size=22, weight="bold"), text_color="#a8e600")
        lbl_total.pack(anchor="w", padx=20)

        txt_recibido.bind("<KeyRelease>", lambda _e: calcular_cambio())

        btn_frame = ctk.CTkFrame(right_panel, fg_color="transparent")
        btn_frame.pack(side="bottom", fill="x", padx=20, pady=20)

        frame_imprimir = ctk.CTkFrame(btn_frame, fg_color="transparent")
        frame_imprimir.pack(anchor="center", pady=(0, 15))
        ctk.CTkLabel(frame_imprimir, text="¿Imprimir Factura?", font=ctk.CTkFont(size=13, weight="bold"), text_color="white").pack(side="left", padx=(0, 10))
        cmb_imprimir = ctk.CTkOptionMenu(frame_imprimir, values=["Sí", "No"], width=90, height=32, font=ctk.CTkFont(size=13))
        cmb_imprimir.pack(side="left")
        cmb_imprimir.set("No")

        def es_mensualidad(nombre_producto):
            nombre = nombre_producto.lower()
            palabras = ("mensual", "membresía", "membresia", "suscripción", "suscripcion")
            return any(p in nombre for p in palabras)

        def cobrar_e_imprimir():
            if not self.carrito_actual:
                messagebox.showerror("Error", "Agregue productos para facturar.")
                return

            # Auto-rellenar el campo "Recibido" con el total del carrito
            total_carrito = sum(item["precio"] * item["cantidad"] for item in self.carrito_actual)
            txt_recibido.delete(0, "end")
            txt_recibido.insert(0, f"{total_carrito:.2f}")
            calcular_cambio()

            cliente_final = cliente_seleccionado["valor"]
            metodo_pago = cmb_metodo_pago.get()
            total = total_carrito

            fecha = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            cliente_es_real = cliente_final and cliente_final.startswith("CLI-") and cliente_final != "CLI-0000 - Consumidor Final"
            contiene_mensualidad = any(es_mensualidad(item["nombre"]) for item in self.carrito_actual)
            cliente_id_renovar = None
            nueva_fecha_renovacion = None

            if cliente_es_real and contiene_mensualidad:
                datos_cliente = clientes_map.get(cliente_final)
                if datos_cliente:
                    cliente_id_renovar = datos_cliente[0]
                    fecha_vencimiento = datos_cliente[4]
                    hoy = datetime.now().date()
                    if fecha_vencimiento:
                        vence = datetime.strptime(fecha_vencimiento, "%Y-%m-%d").date()
                        inicio = max(hoy, vence)
                    else:
                        inicio = hoy
                    cantidad_mensualidades = sum(item["cantidad"] for item in self.carrito_actual if es_mensualidad(item["nombre"]))
                    nueva_fecha_renovacion = (inicio + timedelta(days=30 * cantidad_mensualidades)).strftime("%Y-%m-%d")

            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                conn.execute("BEGIN")
                cursor.execute("INSERT INTO facturas (caja_id, cliente_nombre, total, fecha, metodo_pago) VALUES (?, ?, ?, ?, ?)",
                               (self.id_caja_activa, cliente_final, total, fecha, metodo_pago))
                factura_id = cursor.lastrowid
                cod_fact = f"FACT-{factura_id:05d}"
                cursor.execute("UPDATE facturas SET codigo_factura=? WHERE id=?", (cod_fact, factura_id))
                for item in self.carrito_actual:
                    cursor.execute("INSERT INTO factura_detalles (factura_id, nombre_producto, cantidad, precio_unitario) VALUES (?, ?, ?, ?)",
                                   (factura_id, item["nombre"], item["cantidad"], item["precio"]))
                    cursor.execute("UPDATE productos SET stock=stock-? WHERE id=?", (item["cantidad"], item["id"]))
                    cursor.execute("INSERT INTO historial_inventario (producto_id, nombre_producto, tipo_movimiento, cantidad, fecha, usuario, motivo) VALUES (?, ?, 'SALIDA', ?, ?, ?, ?)",
                                   (item["id"], item["nombre"], item["cantidad"], fecha, self.usuario_actual, f"Venta Factura {cod_fact}"))
                if cliente_id_renovar and nueva_fecha_renovacion:
                    cursor.execute("UPDATE clientes SET fecha_vencimiento=? WHERE id=?", (nueva_fecha_renovacion, cliente_id_renovar))
                conn.commit()
            except Exception as e:
                conn.rollback()
                messagebox.showerror("Error", f"Error al registrar la venta: {e}")
                return
            finally:
                conn.close()

            if cmb_imprimir.get() == "Sí":
                ANCHO = 48
                def fmt(texto, alineacion="left"):
                    if len(texto) >= ANCHO:
                        return texto[:ANCHO]
                    if alineacion == "center":
                        return texto.center(ANCHO)
                    if alineacion == "right":
                        return texto.rjust(ANCHO)
                    return texto.ljust(ANCHO)

                ticket = ""
                ticket += fmt("=" * ANCHO) + "\n"
                ticket += fmt("MONSTER GYM FACTURA", "center") + "\n"
                ticket += fmt("=" * ANCHO) + "\n"
                ticket += fmt(f"Factura: {cod_fact}") + "\n"
                ticket += fmt(f"Cliente: {cliente_final}") + "\n"
                ticket += fmt(f"Fecha: {fecha}") + "\n"
                ticket += fmt(f"Método de pago: {metodo_pago}") + "\n"
                ticket += fmt(f"Atendido por: {self.usuario_actual}") + "\n"
                ticket += fmt("-" * ANCHO) + "\n"
                for item in self.carrito_actual:
                    ticket += fmt(f"{item['nombre']} x{item['cantidad']} : RD${item['precio'] * item['cantidad']:,.2f}") + "\n"
                ticket += fmt("-" * ANCHO) + "\n"
                ticket += fmt(f"TOTAL: RD${total:,.2f}", "right") + "\n"
                if nueva_fecha_renovacion:
                    ticket += fmt(f"Vence: {nueva_fecha_renovacion}") + "\n"
                ticket += fmt("=" * ANCHO) + "\n"

                if self.impresora:
                    try:
                        self.impresora.imprimir(ticket)
                    except Exception as e:
                        messagebox.showerror("Error de impresión", f"No se pudo imprimir:\n{e}")
                        print(ticket)
                else:
                    print(ticket)

            if nueva_fecha_renovacion:
                messagebox.showinfo("Venta completada", f"Venta registrada.\nFactura: {cod_fact}\n\nMensualidad aplicada.\nNuevo vencimiento: {nueva_fecha_renovacion}")
            else:
                messagebox.showinfo("Completado", f"Venta registrada correctamente.\nFactura: {cod_fact}")

            self.carrito_actual = []
            self.mostrar_pantalla_facturacion()

        ctk.CTkButton(btn_frame, text="Cobrar Venta", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=14), height=46,
                      command=cobrar_e_imprimir).pack(fill="x")

        refrescar_tabla_carrito()

    # =========================================================================
    # USUARIOS
    # =========================================================================
    def mostrar_pantalla_usuarios(self):
        self.crear_cabecera("Usuarios", "Administración de cuentas")
        main_split = ctk.CTkFrame(self.workspace, fg_color="transparent")
        main_split.pack(fill="both", expand=True)
        main_split.grid_columnconfigure(0, weight=1)
        main_split.grid_columnconfigure(1, weight=2)
        main_split.grid_rowconfigure(0, weight=1)

        form_frame = ctk.CTkFrame(main_split, fg_color="#14141e", corner_radius=12)
        form_frame.grid(row=0, column=0, padx=8, sticky="nsew")

        ctk.CTkLabel(form_frame, text="Nuevo Usuario", font=ctk.CTkFont(size=16, weight="bold"), text_color="#a8e600").pack(anchor="w", padx=20, pady=(20, 10))
        txt_u = ctk.CTkEntry(form_frame, placeholder_text="Username", height=38, font=ctk.CTkFont(size=14))
        txt_u.pack(fill="x", padx=20, pady=6)
        txt_p = ctk.CTkEntry(form_frame, placeholder_text="Password", show="*", height=38, font=ctk.CTkFont(size=14))
        txt_p.pack(fill="x", padx=20, pady=6)
        cmb_r = ctk.CTkOptionMenu(form_frame, values=["Admin", "Caja"], height=38, font=ctk.CTkFont(size=14))
        cmb_r.pack(fill="x", padx=20, pady=6)

        def crear_usuario():
            try:
                conn = self.conectar_db()
                cursor = conn.cursor()
                cursor.execute("INSERT INTO usuarios (username, password, rol) VALUES (?, ?, ?)", (txt_u.get(), txt_p.get(), cmb_r.get()))
                conn.commit()
                conn.close()
                txt_u.delete(0, "end")
                txt_p.delete(0, "end")
                self.mostrar_pantalla_usuarios()
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo crear: {e}")

        ctk.CTkButton(form_frame, text="Crear Usuario", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=14), height=40, command=crear_usuario).pack(fill="x", padx=20, pady=20)

        list_frame = ctk.CTkScrollableFrame(main_split, fg_color="#14141e", corner_radius=12, height=400)
        list_frame.grid(row=0, column=1, padx=8, sticky="nsew")

        header = ctk.CTkFrame(list_frame, fg_color="#1e1e2a")
        header.pack(fill="x", padx=10, pady=(5, 2))
        ctk.CTkLabel(header, text="ID", width=50, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="USUARIO", width=150, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="ROL", width=80, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="ACCIONES", width=150, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, username, rol FROM usuarios")
        users = cursor.fetchall()
        conn.close()

        if users:
            for id_, u, r in users:
                item = ctk.CTkFrame(list_frame, fg_color="#0d0d14")
                item.pack(fill="x", pady=2, padx=10)
                ctk.CTkLabel(item, text=str(id_), width=50).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=u, width=150).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=r, width=80).pack(side="left", padx=5)
                if self.rol_actual == "Admin":
                    ctk.CTkButton(item, text="Editar", width=60, height=28, fg_color="#1a73e8", command=lambda i=id_: self.editar_usuario(i)).pack(side="left", padx=2)
                    ctk.CTkButton(item, text="Eliminar", width=60, height=28, fg_color="#ff595e", command=lambda i=id_: self.eliminar_usuario(i)).pack(side="left", padx=2)
        else:
            ctk.CTkLabel(list_frame, text="No hay usuarios registrados.", text_color="gray").pack(pady=30)

    def editar_usuario(self, user_id):
        if self.rol_actual != "Admin":
            messagebox.showerror("Error", "No tienes permisos.")
            return
        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT username, rol FROM usuarios WHERE id=?", (user_id,))
        data = cursor.fetchone()
        conn.close()
        if not data:
            return

        ventana = ctk.CTkToplevel(self)
        ventana.title("Editar Usuario")
        ventana.geometry("400x300")
        ventana.grab_set()
        ctk.CTkLabel(ventana, text="Username:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_u = ctk.CTkEntry(ventana, width=320, font=ctk.CTkFont(size=14))
        entry_u.insert(0, data[0])
        entry_u.pack(pady=5)
        ctk.CTkLabel(ventana, text="Nueva Contraseña (opcional):", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        entry_p = ctk.CTkEntry(ventana, width=320, show="*", font=ctk.CTkFont(size=14))
        entry_p.pack(pady=5)
        ctk.CTkLabel(ventana, text="Rol:", font=ctk.CTkFont(weight="bold")).pack(pady=5)
        cmb_r = ctk.CTkOptionMenu(ventana, values=["Admin", "Caja"], font=ctk.CTkFont(size=14))
        cmb_r.set(data[1])
        cmb_r.pack(pady=5)

        def guardar():
            new_user = entry_u.get().strip()
            new_pass = entry_p.get().strip()
            new_rol = cmb_r.get()
            if not new_user:
                messagebox.showerror("Error", "El usuario no puede estar vacío.")
                return
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                if new_pass:
                    cursor.execute("UPDATE usuarios SET username=?, password=?, rol=? WHERE id=?", (new_user, new_pass, new_rol, user_id))
                else:
                    cursor.execute("UPDATE usuarios SET username=?, rol=? WHERE id=?", (new_user, new_rol, user_id))
                conn.commit()
                messagebox.showinfo("Éxito", "Usuario actualizado.")
                ventana.destroy()
                self.mostrar_pantalla_usuarios()
            except sqlite3.IntegrityError:
                messagebox.showerror("Error", "El nombre de usuario ya existe.")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo actualizar: {e}")
            finally:
                conn.close()
        ctk.CTkButton(ventana, text="Guardar Cambios", fg_color="#a8e600", text_color="black", command=guardar).pack(pady=15)

    def eliminar_usuario(self, user_id):
        if self.rol_actual != "Admin":
            messagebox.showerror("Error", "No tienes permisos.")
            return
        if user_id == self.obtener_id_usuario_actual():
            messagebox.showerror("Error", "No puedes eliminar tu propio usuario.")
            return
        if messagebox.askyesno("Confirmar", "¿Eliminar este usuario? Se perderán todos sus datos."):
            conn = self.conectar_db()
            cursor = conn.cursor()
            try:
                cursor.execute("DELETE FROM usuarios WHERE id=?", (user_id,))
                conn.commit()
                messagebox.showinfo("Éxito", "Usuario eliminado.")
                self.mostrar_pantalla_usuarios()
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo eliminar: {e}")
            finally:
                conn.close()

    def obtener_id_usuario_actual(self):
        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM usuarios WHERE username=?", (self.usuario_actual,))
        res = cursor.fetchone()
        conn.close()
        return res[0] if res else None

    # =========================================================================
    # REPORTES (con calendario y exportar a Excel/PDF)
    # =========================================================================
    def mostrar_pantalla_reportes(self):
        self.crear_cabecera("Reportes Avanzados", "Control e histórico general")
        tabview = ctk.CTkTabview(self.workspace, fg_color="#14141e")
        tabview.pack(fill="both", expand=True, padx=10, pady=10)

        tab_ventas_prod = tabview.add("Ventas por Producto")
        tab_cajas = tabview.add("Reporte de Cajas")
        tab_inv = tabview.add("Movimientos de Inventario")

        # ============ TAB VENTAS POR PRODUCTO ============
        filtro_ventas = ctk.CTkFrame(tab_ventas_prod, fg_color="#1a1a26", corner_radius=12)
        filtro_ventas.pack(fill="x", padx=5, pady=5)
        ctk.CTkLabel(filtro_ventas, text="Fecha Desde:", font=ctk.CTkFont(size=13)).pack(side="left", padx=8)
        entry_desde_v = self.crear_selector_fecha(filtro_ventas)
        entry_desde_v.pack(side="left", padx=5)
        ctk.CTkLabel(filtro_ventas, text="Hasta:", font=ctk.CTkFont(size=13)).pack(side="left", padx=8)
        entry_hasta_v = self.crear_selector_fecha(filtro_ventas)
        entry_hasta_v.pack(side="left", padx=5)

        def aplicar_filtro_ventas():
            self.cargar_reporte_ventas_producto(entry_desde_v.get().strip(), entry_hasta_v.get().strip(), report_frame_ventas)

        def exportar_ventas_excel():
            desde = entry_desde_v.get().strip()
            hasta = entry_hasta_v.get().strip()
            query = """
                SELECT fd.nombre_producto, SUM(fd.cantidad), SUM(fd.cantidad * fd.precio_unitario)
                FROM factura_detalles fd JOIN facturas f ON fd.factura_id = f.id WHERE 1=1
            """
            params = []
            if desde: query += " AND DATE(f.fecha) >= ?"; params.append(desde)
            if hasta: query += " AND DATE(f.fecha) <= ?"; params.append(hasta)
            query += " GROUP BY fd.nombre_producto ORDER BY SUM(fd.cantidad) DESC"
            conn = self.conectar_db(); cursor = conn.cursor()
            cursor.execute(query, params); datos = cursor.fetchall(); conn.close()
            columnas = ["Producto", "Cantidad", "Monto (RD$)"]
            self.exportar_reporte_excel(columnas, datos, "Reporte de Ventas por Producto",
                                        f"ventas_producto_{datetime.now().strftime('%Y%m%d')}.xlsx")

        def exportar_ventas_pdf():
            desde = entry_desde_v.get().strip()
            hasta = entry_hasta_v.get().strip()
            query = """
                SELECT fd.nombre_producto, SUM(fd.cantidad), SUM(fd.cantidad * fd.precio_unitario)
                FROM factura_detalles fd JOIN facturas f ON fd.factura_id = f.id WHERE 1=1
            """
            params = []
            if desde: query += " AND DATE(f.fecha) >= ?"; params.append(desde)
            if hasta: query += " AND DATE(f.fecha) <= ?"; params.append(hasta)
            query += " GROUP BY fd.nombre_producto ORDER BY SUM(fd.cantidad) DESC"
            conn = self.conectar_db(); cursor = conn.cursor()
            cursor.execute(query, params); datos = cursor.fetchall(); conn.close()
            columnas = ["Producto", "Cantidad", "Monto (RD$)"]
            self.exportar_reporte_pdf(columnas, datos, "Reporte de Ventas por Producto",
                                      f"ventas_producto_{datetime.now().strftime('%Y%m%d')}.pdf")

        ctk.CTkButton(filtro_ventas, text="🔍 Filtrar", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=12), height=34, command=aplicar_filtro_ventas).pack(side="left", padx=10)
        ctk.CTkButton(filtro_ventas, text="📊 Excel", fg_color="#1a7f37", text_color="white", font=ctk.CTkFont(weight="bold", size=12), height=34, command=exportar_ventas_excel).pack(side="left", padx=3)
        ctk.CTkButton(filtro_ventas, text="📄 PDF", fg_color="#b32121", text_color="white", font=ctk.CTkFont(weight="bold", size=12), height=34, command=exportar_ventas_pdf).pack(side="left", padx=3)

        report_frame_ventas = ctk.CTkScrollableFrame(tab_ventas_prod, fg_color="transparent", corner_radius=12)
        report_frame_ventas.pack(fill="both", expand=True, padx=5, pady=5)
        self.cargar_reporte_ventas_producto("", "", report_frame_ventas)

        # ============ TAB CAJAS ============
        filtro_frame = ctk.CTkFrame(tab_cajas, fg_color="#1a1a26", corner_radius=12)
        filtro_frame.pack(fill="x", padx=5, pady=5)
        ctk.CTkLabel(filtro_frame, text="Fecha Desde:", font=ctk.CTkFont(size=13)).pack(side="left", padx=8)
        entry_desde = self.crear_selector_fecha(filtro_frame)
        entry_desde.pack(side="left", padx=5)
        ctk.CTkLabel(filtro_frame, text="Hasta:", font=ctk.CTkFont(size=13)).pack(side="left", padx=8)
        entry_hasta = self.crear_selector_fecha(filtro_frame)
        entry_hasta.pack(side="left", padx=5)

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT DISTINCT usuario FROM cajas")
        cajeros = [row[0] for row in cursor.fetchall()]
        conn.close()

        ctk.CTkLabel(filtro_frame, text="Cajero:", font=ctk.CTkFont(size=13)).pack(side="left", padx=8)
        cmb_cajero = ctk.CTkOptionMenu(filtro_frame, values=["Todos"] + cajeros, width=140, font=ctk.CTkFont(size=13))
        cmb_cajero.pack(side="left", padx=5)
        cmb_cajero.set("Todos")

        def consultar_cajas():
            desde = entry_desde.get().strip()
            hasta = entry_hasta.get().strip()
            cajero = cmb_cajero.get()
            query = "SELECT id, usuario, monto_apertura, monto_cierre, estado, fecha_apertura FROM cajas WHERE 1=1"
            params = []
            if desde: query += " AND DATE(fecha_apertura) >= ?"; params.append(desde)
            if hasta: query += " AND DATE(fecha_apertura) <= ?"; params.append(hasta)
            if cajero != "Todos" and cajero: query += " AND usuario = ?"; params.append(cajero)
            query += " ORDER BY id DESC"
            conn = self.conectar_db(); cursor = conn.cursor()
            cursor.execute(query, params); datos = cursor.fetchall(); conn.close()
            return datos

        def aplicar_filtros():
            self.cargar_reportes_cajas(entry_desde.get().strip(), entry_hasta.get().strip(), cmb_cajero.get(), report_frame)

        def exportar_cajas_excel():
            datos = consultar_cajas()
            columnas = ["ID Caja", "Cajero", "Apertura (RD$)", "Cierre (RD$)", "Estado", "Fecha Apertura"]
            datos_fmt = [(d[0], d[1], f"{d[2]:,.2f}", f"{d[3]:,.2f}" if d[3] else "—", d[4], d[5]) for d in datos]
            self.exportar_reporte_excel(columnas, datos_fmt, "Reporte de Cajas",
                                        f"reporte_cajas_{datetime.now().strftime('%Y%m%d')}.xlsx")

        def exportar_cajas_pdf():
            datos = consultar_cajas()
            columnas = ["ID", "Cajero", "Apertura", "Cierre", "Estado", "Fecha"]
            datos_fmt = [(d[0], d[1], f"RD${d[2]:,.2f}", f"RD${d[3]:,.2f}" if d[3] else "—", d[4], d[5]) for d in datos]
            self.exportar_reporte_pdf(columnas, datos_fmt, "Reporte de Cajas",
                                      f"reporte_cajas_{datetime.now().strftime('%Y%m%d')}.pdf")

        ctk.CTkButton(filtro_frame, text="🔍 Filtrar", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=12), height=34, command=aplicar_filtros).pack(side="left", padx=10)
        ctk.CTkButton(filtro_frame, text="📊 Excel", fg_color="#1a7f37", text_color="white", font=ctk.CTkFont(weight="bold", size=12), height=34, command=exportar_cajas_excel).pack(side="left", padx=3)
        ctk.CTkButton(filtro_frame, text="📄 PDF", fg_color="#b32121", text_color="white", font=ctk.CTkFont(weight="bold", size=12), height=34, command=exportar_cajas_pdf).pack(side="left", padx=3)

        report_frame = ctk.CTkScrollableFrame(tab_cajas, fg_color="transparent", corner_radius=12)
        report_frame.pack(fill="both", expand=True, padx=5, pady=5)
        self.cargar_reportes_cajas("", "", "Todos", report_frame)

        # ============ TAB INVENTARIO ============
        filtro_inv = ctk.CTkFrame(tab_inv, fg_color="#1a1a26", corner_radius=12)
        filtro_inv.pack(fill="x", padx=5, pady=5)
        ctk.CTkLabel(filtro_inv, text="Fecha Desde:", font=ctk.CTkFont(size=13)).pack(side="left", padx=8)
        entry_desde_inv = self.crear_selector_fecha(filtro_inv)
        entry_desde_inv.pack(side="left", padx=5)
        ctk.CTkLabel(filtro_inv, text="Hasta:", font=ctk.CTkFont(size=13)).pack(side="left", padx=8)
        entry_hasta_inv = self.crear_selector_fecha(filtro_inv)
        entry_hasta_inv.pack(side="left", padx=5)

        def consultar_inventario():
            desde = entry_desde_inv.get().strip()
            hasta = entry_hasta_inv.get().strip()
            query = "SELECT fecha, tipo_movimiento, nombre_producto, cantidad, motivo FROM historial_inventario WHERE 1=1"
            params = []
            if desde: query += " AND DATE(fecha) >= ?"; params.append(desde)
            if hasta: query += " AND DATE(fecha) <= ?"; params.append(hasta)
            query += " ORDER BY id DESC"
            conn = self.conectar_db(); cursor = conn.cursor()
            cursor.execute(query, params); datos = cursor.fetchall(); conn.close()
            return datos

        def aplicar_filtros_inv():
            for widget in scroll_inv.winfo_children():
                widget.destroy()
            datos = consultar_inventario()
            if datos:
                for f, tipo, prod, cant, mot in datos:
                    item = ctk.CTkFrame(scroll_inv, fg_color="#0d0d14")
                    item.pack(fill="x", pady=2, padx=10)
                    color_tipo = "#a8e600" if tipo == "ENTRADA" else "#ff595e"
                    ctk.CTkLabel(item, text=f, width=160).pack(side="left", padx=5)
                    ctk.CTkLabel(item, text=tipo, text_color=color_tipo, width=120, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
                    ctk.CTkLabel(item, text=prod, width=200, anchor="w").pack(side="left", padx=5)
                    ctk.CTkLabel(item, text=str(cant), width=60).pack(side="left", padx=5)
                    ctk.CTkLabel(item, text=mot, width=180, anchor="w").pack(side="left", padx=5)
            else:
                ctk.CTkLabel(scroll_inv, text="No hay registros de movimientos.", text_color="gray").pack(pady=40)

        def exportar_inv_excel():
            datos = consultar_inventario()
            columnas = ["Fecha", "Movimiento", "Producto", "Cantidad", "Motivo"]
            self.exportar_reporte_excel(columnas, datos, "Movimientos de Inventario",
                                        f"inventario_{datetime.now().strftime('%Y%m%d')}.xlsx")

        def exportar_inv_pdf():
            datos = consultar_inventario()
            columnas = ["Fecha", "Movimiento", "Producto", "Cant.", "Motivo"]
            self.exportar_reporte_pdf(columnas, datos, "Movimientos de Inventario",
                                      f"inventario_{datetime.now().strftime('%Y%m%d')}.pdf")

        ctk.CTkButton(filtro_inv, text="🔍 Filtrar", fg_color="#a8e600", text_color="black", font=ctk.CTkFont(weight="bold", size=12), height=34, command=aplicar_filtros_inv).pack(side="left", padx=10)
        ctk.CTkButton(filtro_inv, text="📊 Excel", fg_color="#1a7f37", text_color="white", font=ctk.CTkFont(weight="bold", size=12), height=34, command=exportar_inv_excel).pack(side="left", padx=3)
        ctk.CTkButton(filtro_inv, text="📄 PDF", fg_color="#b32121", text_color="white", font=ctk.CTkFont(weight="bold", size=12), height=34, command=exportar_inv_pdf).pack(side="left", padx=3)

        header_inv = ctk.CTkFrame(tab_inv, fg_color="#1e1e2a")
        header_inv.pack(fill="x", padx=10, pady=(5, 5))
        ctk.CTkLabel(header_inv, text="FECHA", width=160, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header_inv, text="MOVIMIENTO", width=120, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header_inv, text="PRODUCTO", width=200, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header_inv, text="CANT.", width=60, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header_inv, text="MOTIVO", width=180, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)

        scroll_inv = ctk.CTkScrollableFrame(tab_inv, fg_color="transparent")
        scroll_inv.pack(fill="both", expand=True)

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute("SELECT fecha, tipo_movimiento, nombre_producto, cantidad, motivo FROM historial_inventario ORDER BY id DESC")
        movimientos = cursor.fetchall()
        conn.close()

        if movimientos:
            for f, tipo, prod, cant, mot in movimientos:
                item = ctk.CTkFrame(scroll_inv, fg_color="#0d0d14")
                item.pack(fill="x", pady=2, padx=10)
                color_tipo = "#a8e600" if tipo == "ENTRADA" else "#ff595e"
                ctk.CTkLabel(item, text=f, width=160).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=tipo, text_color=color_tipo, width=120, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=prod, width=200, anchor="w").pack(side="left", padx=5)
                ctk.CTkLabel(item, text=str(cant), width=60).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=mot, width=180, anchor="w").pack(side="left", padx=5)
        else:
            ctk.CTkLabel(scroll_inv, text="No hay registros de movimientos en el inventario.", text_color="gray").pack(pady=40)

    def cargar_reporte_ventas_producto(self, desde, hasta, report_frame):
        for widget in report_frame.winfo_children():
            widget.destroy()
        query = """
            SELECT fd.nombre_producto, SUM(fd.cantidad), SUM(fd.cantidad * fd.precio_unitario)
            FROM factura_detalles fd JOIN facturas f ON fd.factura_id = f.id WHERE 1=1
        """
        params = []
        if desde:
            query += " AND DATE(f.fecha) >= ?"
            params.append(desde)
        if hasta:
            query += " AND DATE(f.fecha) <= ?"
            params.append(hasta)
        query += " GROUP BY fd.nombre_producto ORDER BY SUM(fd.cantidad) DESC"

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute(query, params)
        ventas_productos = cursor.fetchall()
        conn.close()

        header = ctk.CTkFrame(report_frame, fg_color="#1e1e2a")
        header.pack(fill="x", padx=10, pady=(5, 2))
        ctk.CTkLabel(header, text="PRODUCTO VENDIDO", width=300, font=ctk.CTkFont(weight="bold"), anchor="w").pack(side="left", padx=20)
        ctk.CTkLabel(header, text="CANTIDAD TOTAL", width=150, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="MONTO GENERADO", width=150, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)

        total_general = 0.0
        if ventas_productos:
            for prod, cant, monto in ventas_productos:
                total_general += monto
                item = ctk.CTkFrame(report_frame, fg_color="#0d0d14")
                item.pack(fill="x", pady=2, padx=10)
                ctk.CTkLabel(item, text=prod, width=300, font=ctk.CTkFont(size=14), anchor="w").pack(side="left", padx=20)
                ctk.CTkLabel(item, text=f"{cant} uds", width=150, font=ctk.CTkFont(size=14)).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=f"RD${monto:,.2f}", width=150, font=ctk.CTkFont(size=14, weight="bold"), text_color="#a8e600").pack(side="left", padx=5)
            frame_total = ctk.CTkFrame(report_frame, fg_color="transparent")
            frame_total.pack(fill="x", pady=20, padx=10)
            ctk.CTkLabel(frame_total, text=f"TOTAL RECAUDADO: RD${total_general:,.2f}", font=ctk.CTkFont(size=18, weight="bold"), text_color="white").pack(side="right", padx=30)
        else:
            ctk.CTkLabel(report_frame, text="No hay ventas registradas en las fechas seleccionadas.", text_color="gray", font=ctk.CTkFont(size=14)).pack(pady=40)

    def cargar_reportes_cajas(self, desde, hasta, cajero, report_frame):
        for widget in report_frame.winfo_children():
            widget.destroy()
        query = "SELECT id, usuario, monto_apertura, monto_cierre, estado, fecha_apertura FROM cajas WHERE 1=1"
        params = []
        if desde:
            query += " AND DATE(fecha_apertura) >= ?"
            params.append(desde)
        if hasta:
            query += " AND DATE(fecha_apertura) <= ?"
            params.append(hasta)
        if cajero != "Todos" and cajero:
            query += " AND usuario = ?"
            params.append(cajero)
        query += " ORDER BY id DESC"

        conn = self.conectar_db()
        cursor = conn.cursor()
        cursor.execute(query, params)
        cajas = cursor.fetchall()
        conn.close()

        header = ctk.CTkFrame(report_frame, fg_color="#1e1e2a")
        header.pack(fill="x", padx=10, pady=(5, 2))
        ctk.CTkLabel(header, text="ID CAJA", width=60, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="CAJERO", width=120, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="APERTURA", width=100, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="CIERRE", width=100, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="ESTADO", width=80, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)
        ctk.CTkLabel(header, text="FECHA APERTURA", width=160, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=5)

        if cajas:
            for c_id, user, m_ap, m_ci, est, f_ap in cajas:
                item = ctk.CTkFrame(report_frame, fg_color="#0d0d14")
                item.pack(fill="x", pady=2, padx=10)
                ctk.CTkLabel(item, text=str(c_id), width=60, font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=user, width=120, font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=f"RD${m_ap:,.2f}", width=100, font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=f"RD${m_ci:,.2f}" if m_ci else "—", width=100, font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=est, width=80, font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
                ctk.CTkLabel(item, text=f_ap, width=160, font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
        else:
            ctk.CTkLabel(report_frame, text="No hay turnos que coincidan con los filtros.", text_color="gray", font=ctk.CTkFont(size=14)).pack(pady=40)


# =============================================================================
if __name__ == "__main__":
    try:
        app = MonsterGymApp()
        app.mainloop()
    except KeyboardInterrupt:
        print("\nSistema cerrado por el usuario.")
    except Exception as e:
        print(f"Error inesperado: {e}")
        messagebox.showerror("Error crítico", f"Ha ocurrido un error:\n{e}")